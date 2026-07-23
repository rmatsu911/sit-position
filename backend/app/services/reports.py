"""帳票生成サービス（共通テーブル帳票基盤）。

すべての帳票は「タイトル＋メタ情報＋表（列定義＋行）」に正規化して生成する。
将来 施工管理表 / 施工写真台帳 / 工程表 / 日報 / 品質チェック表 / 試験記録 / 完成報告書
へ共通化できるよう、build_table_xlsx / build_table_pdf に渡す構造を統一する。
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.platypus.paragraph import Paragraph
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

# 日本語フォント（reportlab同梱のCIDフォント。外部ファイル不要）
_FONT = "HeiseiKakuGo-W5"
try:
    pdfmetrics.registerFont(UnicodeCIDFont(_FONT))
except Exception:  # pragma: no cover
    _FONT = "Helvetica"


class ReportSpec:
    """帳票の正規化構造。"""
    def __init__(self, title: str, meta: list[tuple[str, str]], columns: list[str], rows: list[list[str]]):
        self.title = title
        self.meta = meta            # [(ラベル, 値), ...]
        self.columns = columns      # 列見出し
        self.rows = rows            # 各行の値（文字列）


def build_table_xlsx(spec: ReportSpec) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = spec.title[:28] or "report"

    thin = Side(style="thin", color="B7C2CE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="005BAC")
    header_font = Font(bold=True, color="FFFFFF")

    r = 1
    ws.cell(r, 1, spec.title).font = Font(bold=True, size=14)
    r += 1
    for label, value in spec.meta:
        ws.cell(r, 1, label).font = Font(bold=True)
        ws.cell(r, 2, value)
        r += 1
    r += 1

    header_row = r
    for ci, col in enumerate(spec.columns, start=1):
        c = ws.cell(header_row, ci, col)
        c.fill = header_fill
        c.font = header_font
        c.border = border
        c.alignment = Alignment(horizontal="center", vertical="center")
    r += 1
    for row in spec.rows:
        for ci, val in enumerate(row, start=1):
            c = ws.cell(r, ci, val)
            c.border = border
            c.alignment = Alignment(vertical="center", wrap_text=True)
        r += 1

    for ci, col in enumerate(spec.columns, start=1):
        width = max(len(str(col)), *(len(str(row[ci - 1])) for row in spec.rows)) if spec.rows else len(str(col))
        ws.column_dimensions[get_column_letter(ci)].width = min(40, max(10, width * 1.6))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_table_pdf(spec: ReportSpec) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), topMargin=15 * mm, bottomMargin=15 * mm,
                            leftMargin=12 * mm, rightMargin=12 * mm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("t", parent=styles["Title"], fontName=_FONT, fontSize=16)
    meta_style = ParagraphStyle("m", parent=styles["Normal"], fontName=_FONT, fontSize=9)
    cell_style = ParagraphStyle("c", parent=styles["Normal"], fontName=_FONT, fontSize=8, leading=10)
    head_style = ParagraphStyle("h", parent=styles["Normal"], fontName=_FONT, fontSize=8, leading=10, textColor=colors.white)

    elements = [Paragraph(spec.title, title_style), Spacer(1, 4 * mm)]
    if spec.meta:
        meta_txt = "　".join(f"{k}：{v}" for k, v in spec.meta)
        elements += [Paragraph(meta_txt, meta_style), Spacer(1, 4 * mm)]

    data = [[Paragraph(c, head_style) for c in spec.columns]]
    for row in spec.rows:
        data.append([Paragraph(str(v), cell_style) for v in row])

    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#005BAC")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B7C2CE")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F7FA")]),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    elements.append(table)
    doc.build(elements)
    return buf.getvalue()


def render(spec: ReportSpec, fmt: str) -> tuple[bytes, str]:
    """(bytes, content_type) を返す。"""
    if fmt == "xlsx":
        return build_table_xlsx(spec), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return build_table_pdf(spec), "application/pdf"
