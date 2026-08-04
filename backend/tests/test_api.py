from datetime import timedelta

import pytest

from tests.conftest import token


def test_health(client):
    assert client.get("/health").json()["status"] == "ok"


def test_login_and_me(client):
    t = token(client, "admin@test.jp")
    me = client.get("/api/auth/me", headers={"Authorization": f"Bearer {t}"})
    assert me.status_code == 200
    assert me.json()["role"] == "ADMIN"


def test_login_wrong_password(client):
    r = client.post("/api/auth/login", json={"email": "admin@test.jp", "password": "x"})
    assert r.status_code == 401


def test_projects_requires_auth(client):
    assert client.get("/api/projects").status_code == 401


def test_admin_sees_all_projects(client):
    t = token(client, "admin@test.jp")
    r = client.get("/api/projects", headers={"Authorization": f"Bearer {t}"})
    assert r.status_code == 200
    assert len(r.json()) == 2


def test_partner_scope_limited(client):
    t = token(client, "partner@test.jp")
    r = client.get("/api/projects", headers={"Authorization": f"Bearer {t}"})
    codes = [p["construction_number"] for p in r.json()]
    assert codes == ["T-001"]  # 割当案件のみ


def test_partner_cannot_access_other_project(client):
    t = token(client, "partner@test.jp")
    r = client.get("/api/projects/2", headers={"Authorization": f"Bearer {t}"})
    assert r.status_code == 403


def test_partner_cannot_create_project(client):
    t = token(client, "partner@test.jp")
    r = client.post(
        "/api/projects",
        headers={"Authorization": f"Bearer {t}"},
        json={"construction_number": "X-999", "name": "不可"},
    )
    assert r.status_code == 403


def test_pm_can_create_and_audit(client):
    t = token(client, "pm@test.jp")
    r = client.post(
        "/api/projects",
        headers={"Authorization": f"Bearer {t}"},
        json={"construction_number": "T-100", "name": "新規案件", "customer": "顧客", "status": "未着工"},
    )
    assert r.status_code == 201, r.text
    assert r.json()["construction_number"] == "T-100"


def test_duplicate_construction_number(client):
    t = token(client, "pm@test.jp")
    body = {"construction_number": "T-001", "name": "重複"}
    r = client.post("/api/projects", headers={"Authorization": f"Bearer {t}"}, json=body)
    assert r.status_code == 409


def test_project_create_list_detail_update_reload_flow(client):
    """登録→一覧→詳細→編集→再取得を同じ永続DB経路で保証する。"""
    t = token(client, "pm@test.jp")
    headers = {"Authorization": f"Bearer {t}"}
    created = client.post(
        "/api/projects",
        headers=headers,
        json={
            "construction_number": "KM-2026-E2E",
            "name": "熊本東エリア 光ケーブル更改工事",
            "customer": "株式会社SYSKEN",
            "area": "熊本市東区",
            "status": "未着工",
        },
    )
    assert created.status_code == 201, created.text
    project_id = created.json()["id"]

    listed = client.get("/api/projects", headers=headers)
    assert listed.status_code == 200
    assert any(row["id"] == project_id for row in listed.json())

    detail = client.get(f"/api/projects/{project_id}", headers=headers)
    assert detail.status_code == 200
    assert detail.json()["name"] == "熊本東エリア 光ケーブル更改工事"

    updated = client.put(
        f"/api/projects/{project_id}",
        headers=headers,
        json={"name": "熊本東エリア 光ケーブル更改工事（更新）", "status": "施工中"},
    )
    assert updated.status_code == 200, updated.text

    reloaded = client.get(f"/api/projects/{project_id}", headers=headers)
    assert reloaded.json()["name"].endswith("（更新）")
    assert reloaded.json()["status"] == "施工中"


def test_task_wbs_crud_dependencies_and_reload(client):
    t = token(client, "pm@test.jp")
    headers = {"Authorization": f"Bearer {t}"}
    project_id = 1
    parent = client.post(
        f"/api/projects/{project_id}/tasks",
        headers=headers,
        json={"wbs_code": "90", "name": "E2E親工程", "status": "未着手"},
    )
    assert parent.status_code == 201, parent.text
    parent_id = parent.json()["id"]
    child = client.post(
        f"/api/projects/{project_id}/tasks",
        headers=headers,
        json={
            "parent_task_id": parent_id,
            "wbs_code": "90.1",
            "name": "E2E子工程",
            "planned_start_at": "2026-08-01T00:00:00+09:00",
            "planned_finish_at": "2026-08-02T23:59:59+09:00",
            "planned_workers": 3,
            "dependency_ids": [parent_id],
        },
    )
    assert child.status_code == 201, child.text
    child_id = child.json()["id"]
    assert child.json()["dependencies"] == [parent_id]

    updated = client.put(
        f"/api/tasks/{child_id}",
        headers=headers,
        json={"actual_progress": 60, "actual_workers": 2, "status": "施工中", "change_reason": "E2E"},
    )
    assert updated.status_code == 200, updated.text

    reloaded = client.get(f"/api/projects/{project_id}/tasks", headers=headers)
    row = next(item for item in reloaded.json() if item["id"] == child_id)
    assert row["actual_progress"] == 60
    assert row["actual_workers"] == 2
    assert row["dependencies"] == [parent_id]

    deleted = client.delete(f"/api/tasks/{child_id}", headers=headers)
    assert deleted.status_code == 204
    assert all(item["id"] != child_id for item in client.get(f"/api/projects/{project_id}/tasks", headers=headers).json())


# --- 日時ヘルパー -------------------------------------------------------------
# 本番は PostgreSQL(timestamptz) でタイムゾーン付き、テストは SQLite で
# タイムゾーンを保持できず「JSTの壁時計」がそのまま naive で返る。
# どちらでも同じ意味を検証できるよう、JSTの壁時計(naive)へ正規化する。
def _jst_wall(iso: str):
    from datetime import datetime, timedelta, timezone

    d = datetime.fromisoformat(iso.replace("Z", "+00:00"))
    if d.tzinfo is None:
        return d
    return d.astimezone(timezone(timedelta(hours=9))).replace(tzinfo=None)


def test_half_day_task_roundtrip(client):
    """0.5日(午前/午後)の工程を登録し、再取得しても同じ日時・粒度で返ること。

    期間は [開始, 終了) の半開区間。午前 = 00:00-12:00 / 午後 = 12:00-翌0:00。
    """
    t = token(client, "pm@test.jp")
    headers = {"Authorization": f"Bearer {t}"}
    project_id = 1

    # 午前のみ（0.5日）
    am = client.post(
        f"/api/projects/{project_id}/tasks",
        headers=headers,
        json={
            "wbs_code": "95", "name": "半日工程（午前）",
            "planned_start_at": "2026-06-01T00:00:00+09:00",
            "planned_finish_at": "2026-06-01T12:00:00+09:00",
            "schedule_precision": "half_day",
        },
    )
    assert am.status_code == 201, am.text
    am_id = am.json()["id"]
    assert am.json()["schedule_precision"] == "half_day"

    # 午後開始〜翌日午前終了（1日）
    pm = client.post(
        f"/api/projects/{project_id}/tasks",
        headers=headers,
        json={
            "wbs_code": "96", "name": "半日工程（午後〜翌午前）",
            "planned_start_at": "2026-06-01T12:00:00+09:00",
            "planned_finish_at": "2026-06-02T12:00:00+09:00",
            "schedule_precision": "half_day",
        },
    )
    assert pm.status_code == 201, pm.text

    # 再取得しても同じ値（9時間ずれが起きない）
    rows = client.get(f"/api/projects/{project_id}/tasks", headers=headers).json()
    by_wbs = {r["wbs_code"]: r for r in rows}
    from datetime import timedelta

    a = by_wbs["95"]
    s = _jst_wall(a["planned_start_at"])
    e = _jst_wall(a["planned_finish_at"])
    assert (s.hour, e.hour) == (0, 12), f"午前工程の時刻が想定外: {s} .. {e}"
    assert (e - s) == timedelta(hours=12)  # 0.5日
    assert a["schedule_precision"] == "half_day"

    b = by_wbs["96"]
    s2 = _jst_wall(b["planned_start_at"])
    e2 = _jst_wall(b["planned_finish_at"])
    assert (s2.hour, e2.hour) == (12, 12), f"午後〜翌午前の時刻が想定外: {s2} .. {e2}"
    assert (e2 - s2) == timedelta(days=1)  # 1日

    # 日単位へ変更（区分を持たない工程へ戻せる）
    upd = client.put(
        f"/api/tasks/{am_id}",
        headers=headers,
        json={
            "planned_start_at": "2026-06-01T00:00:00+09:00",
            "planned_finish_at": "2026-06-02T00:00:00+09:00",
            "schedule_precision": "day",
            "change_reason": "日単位へ変更",
        },
    )
    assert upd.status_code == 200, upd.text
    assert upd.json()["schedule_precision"] == "day"
    s3 = _jst_wall(upd.json()["planned_start_at"])
    e3 = _jst_wall(upd.json()["planned_finish_at"])
    assert (e3 - s3) == timedelta(days=1)


def test_day_precision_is_default_and_not_half_day(client):
    """粒度を指定せずに作った工程は日単位（day）になり、半日として解釈されないこと。

    既存データ（Seed）の 09:00 プレースホルダは migration で
    00:00〜翌00:00 の日単位へ正規化済み。ここではAPIの既定動作を検証する。
    """
    t = token(client, "pm@test.jp")
    headers = {"Authorization": f"Bearer {t}"}
    from datetime import timedelta

    created = client.post(
        "/api/projects/1/tasks",
        headers=headers,
        json={
            "wbs_code": "97", "name": "日単位工程",
            "planned_start_at": "2026-06-01T00:00:00+09:00",
            "planned_finish_at": "2026-06-04T00:00:00+09:00",
            # schedule_precision を送らない（既定値の確認）
        },
    )
    assert created.status_code == 201, created.text
    body = created.json()
    assert body["schedule_precision"] == "day", "粒度未指定なら day になること"

    s0 = _jst_wall(body["planned_start_at"])
    e0 = _jst_wall(body["planned_finish_at"])
    assert (s0.hour, s0.minute) == (0, 0), f"開始が 00:00 でない: {s0}"
    assert (e0.hour, e0.minute) == (0, 0), f"終了が 00:00 でない: {e0}"
    span = e0 - s0
    assert span == timedelta(days=3), f"3日間にならない: {span}"
    # 日単位の工程は 12:00 境界を持たない（＝半日として描画されない）
    assert span % timedelta(days=1) == timedelta(0)


# --- Ver.0.3 Phase 2: 横断工程表 ---------------------------------------------
def _cross(client, email: str, qs: str = ""):
    t = token(client, email)
    r = client.get(f"/api/schedule/cross{('?' + qs) if qs else ''}",
                   headers={"Authorization": f"Bearer {t}"})
    assert r.status_code == 200, r.text
    return r.json()


def _make_cross_fixture(client):
    """横断工程表の検証用に、2案件へ工程を作る（担当者・担当会社・未割当を含む）。"""
    t = token(client, "pm@test.jp")
    headers = {"Authorization": f"Bearer {t}"}
    from app.models import Company
    from tests.conftest import TestingSessionLocal

    with TestingSessionLocal() as s:
        names = {c.name: c for c in s.query(Company).all()}
        made = []
        # 同名で別レコードの会社を2つ作る（名前ではなくIDで集約できることの確認用）
        for name in ("横断テスト会社A", "横断テスト会社A", "横断テスト会社B"):
            c = Company(name=name, is_partner=True)
            s.add(c); s.flush(); made.append(c.id)
        s.commit()
    company_a1, company_a2, company_b = made
    pm_id = None
    me = client.get("/api/auth/me", headers=headers)
    if me.status_code == 200:
        pm_id = me.json()["id"]

    created = []
    specs = [
        ("cx-1", "横断 親工程", None, "2026-09-01", "2026-09-11", company_a1, pm_id),
        ("cx-1.1", "横断 子工程A", "cx-1", "2026-09-01", "2026-09-06", company_a1, pm_id),
        ("cx-1.2", "横断 子工程B", "cx-1", "2026-09-07", "2026-09-11", company_a2, None),
        ("cx-2", "横断 未割当工程", None, "2026-09-14", "2026-09-16", None, None),
    ]
    parents: dict[str, int] = {}
    for wbs, name, parent, ps, pe, company, manager in specs:
        body = {
            "wbs_code": wbs, "name": name,
            "planned_start_at": f"{ps}T00:00:00+09:00",
            "planned_finish_at": f"{pe}T00:00:00+09:00",
            "company_id": company, "manager_id": manager,
        }
        if parent:
            body["parent_task_id"] = parents[parent]
        r = client.post("/api/projects/2/tasks", headers=headers, json=body)
        assert r.status_code == 201, r.text
        parents[wbs] = r.json()["id"]
        created.append(r.json())
    return {"ids": parents, "company_a1": company_a1, "company_a2": company_a2,
            "company_b": company_b, "manager_id": pm_id}


def test_cross_schedule_groups_by_id_not_name(client):
    """同名の担当会社でも、IDが違えば別グループとして集約されること。"""
    fx = _make_cross_fixture(client)
    a1 = _cross(client, "admin@test.jp", f"company_ids={fx['company_a1']}")
    a2 = _cross(client, "admin@test.jp", f"company_ids={fx['company_a2']}")
    names1 = {t["name"] for t in a1["tasks"] if t["matched"]}
    names2 = {t["name"] for t in a2["tasks"] if t["matched"]}
    assert "横断 子工程A" in names1 and "横断 子工程B" not in names1
    assert "横断 子工程B" in names2 and "横断 子工程A" not in names2
    # 会社名は同じでも別IDとして返る
    assert {t["company"] for t in a1["tasks"] if t["matched"]} == {"横断テスト会社A"}
    assert {t["company"] for t in a2["tasks"] if t["matched"]} == {"横断テスト会社A"}


def test_cross_schedule_keeps_wbs_parent_when_child_matches(client):
    """子工程だけが条件に一致しても、親工程が matched=False で補われること。"""
    fx = _make_cross_fixture(client)
    data = _cross(client, "admin@test.jp", f"company_ids={fx['company_a2']}")
    by_name = {t["name"]: t for t in data["tasks"]}
    assert by_name["横断 子工程B"]["matched"] is True
    assert "横断 親工程" in by_name, "WBSの親工程が補われていない"
    assert by_name["横断 親工程"]["matched"] is False
    assert data["total"] == sum(1 for t in data["tasks"] if t["matched"])


def test_cross_schedule_includes_unassigned(client):
    """未割当の工程を除外せず、未割当のみの絞り込みもできること。"""
    _make_cross_fixture(client)
    allrows = _cross(client, "admin@test.jp")
    assert any(t["manager_id"] is None and t["company_id"] is None for t in allrows["tasks"])

    only = _cross(client, "admin@test.jp", "unassigned_only=true")
    matched = [t for t in only["tasks"] if t["matched"]]
    assert matched, "未割当のみの結果が空"
    assert all(t["manager_id"] is None or t["company_id"] is None for t in matched)


def test_cross_schedule_applies_project_scope(client):
    """案件スコープ外の工程はAPIが返さない（UIの出し分けに依存しない）。"""
    _make_cross_fixture(client)
    partner = _cross(client, "partner@test.jp")
    assert partner["tasks"], "割当案件の工程が取得できない"
    assert {t["project_id"] for t in partner["tasks"]} == {1}

    # フィルタでスコープ外の案件を指定しても返らない
    filtered = _cross(client, "partner@test.jp", "project_ids=2")
    assert filtered["total"] == 0
    assert filtered["tasks"] == []

    # 単一案件表示でもAPI側で拒否する
    t = token(client, "partner@test.jp")
    r = client.get("/api/schedule/cross?project_id=2", headers={"Authorization": f"Bearer {t}"})
    assert r.status_code == 403


def test_cross_schedule_options_are_scoped(client):
    """絞り込みの選択肢も権限と案件スコープの範囲だけを返すこと。"""
    _make_cross_fixture(client)
    t = token(client, "partner@test.jp")
    r = client.get("/api/schedule/cross/options", headers={"Authorization": f"Bearer {t}"})
    assert r.status_code == 200, r.text
    o = r.json()
    assert [p["construction_number"] for p in o["projects"]] == ["T-001"]


def test_cross_schedule_half_day_roundtrip(client):
    """0.5日の工程が横断工程表でも 12:00 境界のまま返ること（9時間ずれ防止）。"""
    from datetime import timedelta

    t = token(client, "pm@test.jp")
    headers = {"Authorization": f"Bearer {t}"}
    r = client.post("/api/projects/2/tasks", headers=headers, json={
        "wbs_code": "cx-9", "name": "横断 半日工程",
        "planned_start_at": "2026-09-21T12:00:00+09:00",
        "planned_finish_at": "2026-09-22T00:00:00+09:00",
        "schedule_precision": "half_day",
    })
    assert r.status_code == 201, r.text

    data = _cross(client, "admin@test.jp", "q=横断 半日工程")
    row = next(x for x in data["tasks"] if x["name"] == "横断 半日工程")
    s = _jst_wall(row["planned_start_at"])
    e = _jst_wall(row["planned_finish_at"])
    assert (s.hour, e.hour) == (12, 0)
    assert (e - s) == timedelta(hours=12)
    assert row["schedule_precision"] == "half_day"


def test_cross_schedule_detections_are_not_ai(client):
    """システム検知は確定計算のみ。AI予測の語や信頼度を含まないこと。"""
    _make_cross_fixture(client)
    data = _cross(client, "admin@test.jp")
    kinds = {d["kind"] for d in data["detections"]}
    assert "unassigned" in kinds
    for d in data["detections"]:
        assert "AI" not in d["label"] and "AI" not in d["message"]
        assert "confidence" not in d
        assert d["task_ids"], "検知に対象工程が無い"


def test_cross_schedule_range_is_derived_from_tasks(client):
    """表示期間は固定日付ではなく、対象工程の実期間から返ること。"""
    _make_cross_fixture(client)
    data = _cross(client, "admin@test.jp", "q=横断 未割当工程")
    row = next(x for x in data["tasks"] if x["matched"])
    assert _jst_wall(data["range_from"]) == _jst_wall(row["planned_start_at"])
    assert _jst_wall(data["range_to"]) == _jst_wall(row["planned_finish_at"])


def test_cross_schedule_export_reflects_filters(client):
    """Excel/PDF 出力が現在の絞り込み条件で生成され、スコープも適用されること。"""
    _make_cross_fixture(client)
    t = token(client, "admin@test.jp")
    headers = {"Authorization": f"Bearer {t}"}
    for fmt in ("xlsx", "pdf"):
        r = client.get(f"/api/schedule/cross/export?format={fmt}&unassigned_only=true", headers=headers)
        assert r.status_code == 200, r.text
        assert len(r.content) > 500
        assert "attachment;" in r.headers["content-disposition"]

    tp = token(client, "partner@test.jp")
    r = client.get("/api/schedule/cross/export?format=xlsx&project_ids=2",
                   headers={"Authorization": f"Bearer {tp}"})
    assert r.status_code == 200  # 生成はできるが中身はスコープ内のみ


def test_saved_search_is_stored_per_user(client):
    """保存検索条件はDBに保存され、他ユーザーからは見えないこと。"""
    ta = {"Authorization": f"Bearer {token(client, 'admin@test.jp')}"}
    tp = {"Authorization": f"Bearer {token(client, 'partner@test.jp')}"}
    conditions = {"delayed_only": True, "statuses": ["遅延"], "scale": "week"}
    created = client.post("/api/schedule/saved-searches", headers=ta,
                          json={"screen": "cross_schedule", "name": "遅延のみ", "conditions": conditions})
    assert created.status_code == 201, created.text
    sid = created.json()["id"]
    assert created.json()["conditions"] == conditions

    assert any(r["id"] == sid for r in client.get("/api/schedule/saved-searches", headers=ta).json())
    assert all(r["id"] != sid for r in client.get("/api/schedule/saved-searches", headers=tp).json())
    assert client.delete(f"/api/schedule/saved-searches/{sid}", headers=tp).status_code == 404
    assert client.delete(f"/api/schedule/saved-searches/{sid}", headers=ta).status_code == 204


def test_cross_schedule_update_uses_ids(client):
    """担当者・担当会社の変更が既存の工程更新APIでID指定のまま保存されること。"""
    fx = _make_cross_fixture(client)
    task_id = fx["ids"]["cx-2"]
    headers = {"Authorization": f"Bearer {token(client, 'pm@test.jp')}"}
    r = client.put(f"/api/tasks/{task_id}", headers=headers,
                   json={"company_id": fx["company_b"], "change_reason": "担当会社変更"})
    assert r.status_code == 200, r.text
    assert r.json()["company_id"] == fx["company_b"]
    assert r.json()["company"] == "横断テスト会社B"

    # 権限が無いユーザーはAPI側で拒否される
    denied = client.put(f"/api/tasks/{task_id}",
                        headers={"Authorization": f"Bearer {token(client, 'partner@test.jp')}"},
                        json={"actual_progress": 50})
    assert denied.status_code == 403


# --- Ver.0.3 Phase 3 (P3-1): マイルストーンのデータ構造 ---------------------
def test_milestone_is_stored_as_record_not_task_name(client):
    """マイルストーンが工程名の文字列一致ではなく、実レコードとして保持されること。"""
    from app.models import Milestone, MilestoneType, Project
    from tests.conftest import TestingSessionLocal
    from datetime import datetime, timedelta

    with TestingSessionLocal() as s:
        mt = MilestoneType(code="ms-引き渡し", name="引き渡し", sort_order=4)
        s.add(mt)
        s.flush()
        project = s.query(Project).filter_by(construction_number="T-001").one()
        s.add(Milestone(
            project_id=project.id, milestone_type_id=mt.id, name="引き渡し（テスト案件1）",
            planned_at=datetime.fromisoformat("2026-08-10T00:00:00+09:00"),
            actual_at=datetime.fromisoformat("2026-08-13T00:00:00+09:00"),
            status="完了",
        ))
        s.commit()

        row = s.query(Milestone).filter_by(name="引き渡し（テスト案件1）").one()
        assert row.milestone_type_id == mt.id, "区分がIDで紐づくこと（工程名に依存しない）"
        assert row.project_id == project.id
        # 予定・実績はどちらも JST のその日 00:00 として保持する
        assert _jst_wall(row.planned_at.isoformat()).hour == 0
        assert _jst_wall(row.actual_at.isoformat()).hour == 0
        # 遅延は列ではなく確定計算で求める（列を持たないことの確認）
        assert not hasattr(row, "is_delayed")
        delay = _jst_wall(row.actual_at.isoformat()) - _jst_wall(row.planned_at.isoformat())
        assert delay == timedelta(days=3), f"遅れ日数が計算できること: {delay}"


def test_milestone_supports_multiple_types_per_project(client):
    """1案件に複数区分の重要日を持てること（引き渡し以外も扱える）。"""
    from app.models import Milestone, MilestoneType, Project
    from tests.conftest import TestingSessionLocal
    from datetime import datetime

    with TestingSessionLocal() as s:
        project = s.query(Project).filter_by(construction_number="T-002").one()
        made = []
        for i, name in enumerate(["契約", "着工", "中間検査", "完成検査", "完工"]):
            mt = MilestoneType(code=f"ms2-{name}", name=name, sort_order=i)
            s.add(mt)
            s.flush()
            made.append(mt)
            s.add(Milestone(
                project_id=project.id, milestone_type_id=mt.id, name=f"{name}（テスト案件2）",
                planned_at=datetime.fromisoformat(f"2026-09-{10 + i:02d}T00:00:00+09:00"),
                status="予定",
            ))
        s.commit()

        rows = (
            s.query(Milestone)
            .filter(Milestone.project_id == project.id, Milestone.deleted_at.is_(None))
            .all()
        )
        assert len(rows) == 5, "区分ごとに複数の重要日を持てること"
        assert len({r.milestone_type_id for r in rows}) == 5, "区分がそれぞれ別レコードであること"
        # 実績未入力は未入力のまま（固定値で埋めない）
        assert all(r.actual_at is None for r in rows)
        assert all(r.status == "予定" for r in rows)


# --- Ver.0.3 Phase 3 (P3-1補足): 担当会社 / 関連工程 / 入力粒度 ---------------
def _mk_milestone(session, **kw):
    from app.models import Milestone, MilestoneType

    mt = session.query(MilestoneType).first()
    if mt is None:
        mt = MilestoneType(code="ms-sup", name="引き渡し", sort_order=0)
        session.add(mt)
        session.flush()
    base = dict(project_id=1, milestone_type_id=mt.id, name="補足テスト", status="予定")
    base.update(kw)
    row = Milestone(**base)
    session.add(row)
    session.flush()
    return row


def test_milestone_company_is_independent_from_responsible(client):
    """担当会社が担当者の所属からの導出ではなく、独立した正データとして保存・再取得できること。"""
    from app.models import Company, Milestone, User
    from tests.conftest import TestingSessionLocal

    with TestingSessionLocal() as s:
        company = Company(name="マイルストーン担当会社", is_partner=True)
        s.add(company)
        s.flush()
        # 担当者には別の所属会社を持たせ、導出されていないことを確かめる
        other = Company(name="担当者の所属会社", is_partner=False)
        s.add(other)
        s.flush()
        user = s.query(User).filter_by(email="pm@test.jp").one()
        user.company_id = other.id
        row = _mk_milestone(s, responsible_id=user.id, company_id=company.id)
        s.commit()
        # セッションを閉じた後に参照しないよう、IDを値として取り出しておく
        mid, company_id, other_id, user_id = row.id, company.id, other.id, user.id

    with TestingSessionLocal() as s:
        got = s.get(Milestone, mid)
        assert got.company_id == company_id, "担当会社がそのまま再取得できること"
        assert got.company_id != other_id, "担当者の所属会社から導出していないこと"
        assert got.responsible_id == user_id


def test_milestone_related_task_roundtrip_and_validation(client):
    """関連工程の保存・再取得と、不正な工程IDを拒否すること。"""
    from fastapi import HTTPException

    from app.models import Milestone, Task, User
    from app.services.milestones import resolve_related_task
    from tests.conftest import TestingSessionLocal

    with TestingSessionLocal() as s:
        pm = s.query(User).filter_by(email="pm@test.jp").one()
        partner = s.query(User).filter_by(email="partner@test.jp").one()

        t1 = Task(project_id=1, wbs_code="ms-1", name="案件1の工程")
        t2 = Task(project_id=2, wbs_code="ms-2", name="案件2の工程")
        gone = Task(project_id=1, wbs_code="ms-3", name="削除済み工程")
        s.add_all([t1, t2, gone])
        s.flush()
        from datetime import datetime, timezone as _tz
        gone.deleted_at = datetime.now(_tz.utc)
        s.flush()

        # 同じ案件の工程は通る
        assert resolve_related_task(s, pm, 1, t1.id) == t1.id
        # 関連付けなしは推測せず None のまま
        assert resolve_related_task(s, pm, 1, None) is None

        # 別案件の工程IDは拒否
        with pytest.raises(HTTPException) as e1:
            resolve_related_task(s, pm, 1, t2.id)
        assert e1.value.status_code == 422

        # 存在しない工程IDは拒否
        with pytest.raises(HTTPException) as e2:
            resolve_related_task(s, pm, 1, 999999)
        assert e2.value.status_code == 422

        # 削除済み工程は拒否
        with pytest.raises(HTTPException) as e3:
            resolve_related_task(s, pm, 1, gone.id)
        assert e3.value.status_code == 422

        # 案件スコープ外の工程は拒否（協力会社ユーザーは案件1のみ）
        with pytest.raises(HTTPException) as e4:
            resolve_related_task(s, partner, 2, t2.id)
        assert e4.value.status_code == 403

        row = _mk_milestone(s, related_task_id=t1.id)
        s.commit()
        mid, tid = row.id, t1.id

    with TestingSessionLocal() as s:
        got = s.get(Milestone, mid)
        assert got.related_task_id == tid, "関連工程がそのまま再取得できること"


def test_milestone_precision_roundtrip_keeps_jst(client):
    """day / half_day（午前・午後）が往復し、9時間ずれないこと。"""
    from app.models import Milestone
    from app.services.milestones import (
        milestone_at, normalize_precision, split_milestone_at,
    )
    from tests.conftest import TestingSessionLocal

    with TestingSessionLocal() as s:
        day = _mk_milestone(s, planned_at=milestone_at("2026-08-10", "day"), schedule_precision="day")
        am = _mk_milestone(s, planned_at=milestone_at("2026-08-11", "half_day", "AM"),
                           schedule_precision="half_day")
        pm = _mk_milestone(s, planned_at=milestone_at("2026-08-12", "half_day", "PM"),
                           schedule_precision="half_day")
        s.commit()
        ids = (day.id, am.id, pm.id)

    with TestingSessionLocal() as s:
        got_day, got_am, got_pm = (s.get(Milestone, i) for i in ids)

        # day は JST 00:00
        assert _jst_wall(got_day.planned_at.isoformat()).hour == 0
        assert got_day.schedule_precision == "day"
        assert split_milestone_at(got_day.planned_at) == ("2026-08-10", "AM")

        # half_day 午前は JST 00:00
        assert _jst_wall(got_am.planned_at.isoformat()).hour == 0
        assert split_milestone_at(got_am.planned_at) == ("2026-08-11", "AM")

        # half_day 午後は JST 12:00
        assert _jst_wall(got_pm.planned_at.isoformat()).hour == 12, "午後は12:00で保存されること"
        assert split_milestone_at(got_pm.planned_at) == ("2026-08-12", "PM")
        assert got_pm.schedule_precision == "half_day"

        # 日付が前後の日へずれていない（9時間ずれの検出）
        for row, expect in ((got_day, "2026-08-10"), (got_am, "2026-08-11"), (got_pm, "2026-08-12")):
            assert split_milestone_at(row.planned_at)[0] == expect, f"9時間ずれ: {row.planned_at}"

    # 想定外の粒度は day として扱う
    assert normalize_precision(None) == "day"
    assert normalize_precision("time") == "day"
    assert normalize_precision("half_day") == "half_day"


def test_milestone_has_no_am_pm_flag_column(client):
    """AM/PM専用の正データ列を作っていないこと（日時とフラグの二重管理をしない）。"""
    from app.models import Milestone

    names = {c.name for c in Milestone.__table__.columns}
    forbidden = {"half", "half_day", "am_pm", "ampm", "is_pm", "is_afternoon", "period", "half_flag"}
    assert not (names & forbidden), f"AM/PM専用列がある: {names & forbidden}"
    # 粒度は schedule_precision の1列だけで表す
    assert "schedule_precision" in names
    assert {"planned_at", "actual_at", "company_id", "related_task_id"} <= names


# --- Ver.0.3 Phase 3 (P3-2): 横断マイルストーンの集約API ---------------------
def _ms_fixture(session, *, project_id=1, type_name="引き渡し", order=4, **kw):
    """テスト用のマイルストーンを1件作る（種別はIDで紐づける）。"""
    from app.models import Milestone, MilestoneType

    mt = session.query(MilestoneType).filter_by(code=f"p32-{type_name}-{order}").one_or_none()
    if mt is None:
        mt = MilestoneType(code=f"p32-{type_name}-{order}", name=type_name, sort_order=order)
        session.add(mt)
        session.flush()
    base = dict(project_id=project_id, milestone_type_id=mt.id, name=f"{type_name}#{order}",
                status="予定", schedule_precision="day")
    base.update(kw)
    row = Milestone(**base)
    session.add(row)
    session.flush()
    return row, mt


def _ms_get(client, email, qs=""):
    t = token(client, email)
    r = client.get(f"/api/schedule/milestones{('?' + qs) if qs else ''}",
                   headers={"Authorization": f"Bearer {t}"})
    assert r.status_code == 200, r.text
    return r.json()


def test_milestone_api_facts_use_jst_today(client):
    """確定計算が JST の基準日で行われること（当日は期限超過にしない）。"""
    from app.services.milestones import jst_today, milestone_at
    from tests.conftest import TestingSessionLocal

    today = jst_today()
    with TestingSessionLocal() as s:
        past, _ = _ms_fixture(s, order=90, type_name="期限超過",
                              planned_at=milestone_at((today - timedelta(days=5)).strftime("%Y-%m-%d")))
        same, _ = _ms_fixture(s, order=91, type_name="当日",
                              planned_at=milestone_at(today.strftime("%Y-%m-%d")))
        soon, _ = _ms_fixture(s, order=92, type_name="近日",
                              planned_at=milestone_at((today + timedelta(days=3)).strftime("%Y-%m-%d")))
        far, _ = _ms_fixture(s, order=93, type_name="先",
                             planned_at=milestone_at((today + timedelta(days=60)).strftime("%Y-%m-%d")))
        done, _ = _ms_fixture(
            s, order=94, type_name="完了遅れ",
            planned_at=milestone_at((today - timedelta(days=10)).strftime("%Y-%m-%d")),
            actual_at=milestone_at((today - timedelta(days=6)).strftime("%Y-%m-%d")), status="完了")
        s.commit()
        ids = {"past": past.id, "same": same.id, "soon": soon.id, "far": far.id, "done": done.id}

    data = _ms_get(client, "admin@test.jp", "due_soon_days=7&include_candidates=false&limit=5000")
    by_id = {m["id"]: m for m in data["milestones"]}
    assert data["calculated_at"] == today.date().isoformat(), "基準日はサーバー側のJSTの日付"

    p = by_id[ids["past"]]
    assert p["is_overdue"] and p["delay_days"] == 5 and p["remaining_days"] is None
    assert p["actual_missing"] and not p["is_completed"]

    sm = by_id[ids["same"]]
    assert not sm["is_overdue"], "当日は期限超過にしない"
    assert sm["remaining_days"] == 0 and sm["is_due_soon"]

    so = by_id[ids["soon"]]
    assert so["is_due_soon"] and so["remaining_days"] == 3 and so["delay_days"] == 0

    fa = by_id[ids["far"]]
    assert not fa["is_due_soon"] and fa["remaining_days"] == 60

    dn = by_id[ids["done"]]
    assert dn["is_completed"] and dn["was_delayed"] and dn["delay_days"] == 4
    assert not dn["is_overdue"] and dn["remaining_days"] is None
    # DBの status と確定計算を混同しない
    assert dn["status"] == "完了" and "is_delayed" not in dn


def test_milestone_api_filters_are_and_combined(client):
    """複合フィルターがANDで適用されること。"""
    from app.models import Company, User
    from app.services.milestones import jst_today, milestone_at
    from tests.conftest import TestingSessionLocal

    today = jst_today()
    with TestingSessionLocal() as s:
        co = Company(name="MS絞り込み会社", is_partner=True)
        s.add(co)
        s.flush()
        pm = s.query(User).filter_by(email="pm@test.jp").one()
        hit, mt = _ms_fixture(
            s, order=95, type_name="AND対象", name="AND対象マイルストーン",
            planned_at=milestone_at((today + timedelta(days=2)).strftime("%Y-%m-%d")),
            company_id=co.id, responsible_id=pm.id, status="予定")
        # 会社だけ違う（AND なので落ちる）
        _ms_fixture(s, order=95, type_name="AND対象", name="別会社",
                    planned_at=milestone_at((today + timedelta(days=2)).strftime("%Y-%m-%d")),
                    responsible_id=pm.id, status="予定")
        s.commit()
        hit_id, type_id, company_id, user_id = hit.id, mt.id, co.id, pm.id

    qs = (f"milestone_type_ids={type_id}&company_ids={company_id}&responsible_ids={user_id}"
          f"&statuses=予定&actual=missing&due_soon_only=true&q=AND対象"
          f"&include_candidates=false&limit=5000")
    data = _ms_get(client, "admin@test.jp", qs)
    ids = [m["id"] for m in data["milestones"]]
    assert ids == [hit_id], f"AND適用の結果が想定と違う: {ids}"
    assert data["registered_count"] == 1

    # 条件を1つでも外すと一致しない（ANDであることの裏取り）
    other = _ms_get(client, "admin@test.jp", f"company_ids={company_id}&statuses=完了&include_candidates=false")
    assert other["registered_count"] == 0


def test_milestone_api_scope_and_invalid_ids(client):
    """案件スコープの適用と、不正ID・スコープ外IDで検索範囲が広がらないこと。"""
    from app.services.milestones import jst_today, milestone_at
    from tests.conftest import TestingSessionLocal

    today = jst_today()
    with TestingSessionLocal() as s:
        _ms_fixture(s, project_id=1, order=96, type_name="案件1のみ",
                    planned_at=milestone_at(today.strftime("%Y-%m-%d")))
        _ms_fixture(s, project_id=2, order=97, type_name="案件2のみ",
                    planned_at=milestone_at(today.strftime("%Y-%m-%d")))
        s.commit()

    # 協力会社ユーザーは案件1のみ
    partner = _ms_get(client, "partner@test.jp", "include_candidates=false&limit=5000")
    assert {m["project_id"] for m in partner["milestones"]} == {1}

    # スコープ外の案件IDを指定しても範囲は広がらない
    filtered = _ms_get(client, "partner@test.jp", "project_ids=2&include_candidates=false")
    assert filtered["registered_count"] == 0
    assert filtered["milestones"] == []

    # 単一案件表示はAPI側で拒否する
    t = token(client, "partner@test.jp")
    r = client.get("/api/schedule/milestones?project_id=2", headers={"Authorization": f"Bearer {t}"})
    assert r.status_code == 403

    # 存在しないIDを渡しても件数は増えない（無視されるだけ）
    bogus = _ms_get(client, "partner@test.jp", "company_ids=999999&responsible_ids=999999&include_candidates=false")
    assert bogus["registered_count"] == 0

    # 選択肢にもスコープが効く
    opts = client.get("/api/schedule/milestones/options",
                      headers={"Authorization": f"Bearer {t}"}).json()
    assert [p["construction_number"] for p in opts["projects"]] == ["T-001"]


def test_milestone_api_excludes_soft_deleted(client):
    """論理削除したマイルストーンと、削除済み関連工程を返さないこと。"""
    from datetime import datetime, timezone as _tz

    from app.models import Milestone, Task
    from app.services.milestones import jst_today, milestone_at
    from tests.conftest import TestingSessionLocal

    today = jst_today()
    with TestingSessionLocal() as s:
        task = Task(project_id=1, wbs_code="ms-del", name="あとで削除する工程",
                    planned_start_at=milestone_at(today.strftime("%Y-%m-%d")),
                    planned_finish_at=milestone_at((today + timedelta(days=2)).strftime("%Y-%m-%d")))
        s.add(task)
        s.flush()
        alive, _ = _ms_fixture(s, order=98, type_name="生存",
                               planned_at=milestone_at(today.strftime("%Y-%m-%d")),
                               related_task_id=task.id)
        dead, _ = _ms_fixture(s, order=99, type_name="削除済み",
                              planned_at=milestone_at(today.strftime("%Y-%m-%d")))
        dead.deleted_at = datetime.now(_tz.utc)
        s.commit()
        alive_id, dead_id, task_id = alive.id, dead.id, task.id

    data = _ms_get(client, "admin@test.jp", "include_candidates=false&limit=5000")
    ids = {m["id"] for m in data["milestones"]}
    assert alive_id in ids and dead_id not in ids, "論理削除したマイルストーンを除外すること"
    got = next(m for m in data["milestones"] if m["id"] == alive_id)
    assert got["related_task_id"] == task_id and got["related_task_name"] == "あとで削除する工程"

    # 関連工程を論理削除すると、関連工程としては返さない
    with TestingSessionLocal() as s:
        s.get(Task, task_id).deleted_at = datetime.now(_tz.utc)
        s.commit()
    after = _ms_get(client, "admin@test.jp", "include_candidates=false&limit=5000")
    got2 = next(m for m in after["milestones"] if m["id"] == alive_id)
    assert got2["related_task_id"] is None and got2["related_task_name"] is None
    opts = client.get("/api/schedule/milestones/options",
                      headers={"Authorization": f"Bearer {token(client, 'admin@test.jp')}"}).json()
    assert all(t["id"] != task_id for t in opts["related_tasks"]), "削除済み工程を選択肢に出さない"


def test_milestone_api_candidates_are_separated(client):
    """未設定候補が登録済みレコードと別配列で返り、架空の値を持たないこと。"""
    from app.models import MilestoneType
    from app.services.milestones import jst_today, milestone_at
    from tests.conftest import TestingSessionLocal

    today = jst_today()
    with TestingSessionLocal() as s:
        registered = MilestoneType(code="cand-登録済み", name="登録済み種別", sort_order=200)
        missing = MilestoneType(code="cand-未設定", name="未設定種別", sort_order=201)
        inactive = MilestoneType(code="cand-非active", name="非active種別", sort_order=202, active=False)
        s.add_all([registered, missing, inactive])
        s.flush()
        from app.models import Milestone
        s.add(Milestone(project_id=1, milestone_type_id=registered.id, name="登録済み",
                        planned_at=milestone_at(today.strftime("%Y-%m-%d")), status="予定"))
        s.commit()
        reg_id, miss_id, inact_id = registered.id, missing.id, inactive.id

    data = _ms_get(client, "admin@test.jp", "project_ids=1&limit=5000")
    cand_type_ids = {c["milestone_type_id"] for c in data["candidates"]}
    ms_type_ids = {m["milestone_type_id"] for m in data["milestones"]}

    assert reg_id in ms_type_ids, "登録済みは milestones 側に入る"
    assert reg_id not in cand_type_ids, "登録済みを候補にしない（IDで比較）"
    assert miss_id in cand_type_ids, "未登録の active 種別は候補になる"
    assert inact_id not in cand_type_ids, "非active の種別を候補にしない"

    for c in data["candidates"]:
        assert c["record_kind"] == "candidate"
        # 実在するID・架空の予定日・実績日・状態・担当者・担当会社を持たない
        assert "id" not in c and "planned_at" not in c and "actual_at" not in c
        assert "status" not in c and "responsible_id" not in c and "company_id" not in c
    assert data["candidate_count"] == len(data["candidates"])
    assert data["registered_count"] == len(data["milestones"])
    assert data["candidate_count"] != data["registered_count"] or data["candidate_count"] == 0


def test_milestone_api_limit_and_sort_are_stable(client):
    """total / returned_count / truncated と、並び順が安定していること。"""
    first = _ms_get(client, "admin@test.jp", "include_candidates=false&limit=3")
    assert first["returned_count"] == len(first["milestones"]) == 3
    assert first["total"] >= first["returned_count"]
    assert first["truncated"] is (first["total"] > first["returned_count"])
    # 絞り込み前件数と混同しない
    assert first["registered_count"] == first["total"]

    order1 = [m["id"] for m in _ms_get(client, "admin@test.jp", "include_candidates=false&limit=5000")["milestones"]]
    order2 = [m["id"] for m in _ms_get(client, "admin@test.jp", "include_candidates=false&limit=5000")["milestones"]]
    assert order1 == order2, "同じ条件なら常に同じ順序"
    assert order1[:3] == [m["id"] for m in first["milestones"]], "上限をかけても先頭は同じ"


def test_milestone_api_same_name_separated_by_id(client):
    """同名のマイルストーンでもIDで分離されること。"""
    from app.services.milestones import jst_today, milestone_at
    from tests.conftest import TestingSessionLocal

    today = jst_today()
    with TestingSessionLocal() as s:
        a, _ = _ms_fixture(s, project_id=1, order=210, type_name="同名", name="同じ名前の重要日",
                           planned_at=milestone_at(today.strftime("%Y-%m-%d")))
        b, _ = _ms_fixture(s, project_id=2, order=211, type_name="同名", name="同じ名前の重要日",
                           planned_at=milestone_at(today.strftime("%Y-%m-%d")))
        s.commit()
        ids = {a.id, b.id}

    data = _ms_get(client, "admin@test.jp", "q=同じ名前の重要日&include_candidates=false&limit=5000")
    got = [m for m in data["milestones"] if m["id"] in ids]
    assert len(got) == 2, "同名でも別レコードとして返る"
    assert {m["project_id"] for m in got} == {1, 2}
    assert len({m["id"] for m in got}) == 2


def test_milestone_api_precision_and_related_task_project_match(client):
    """day / half_day 午前・午後が返り、関連工程が同じ案件であること。"""
    from app.models import Task
    from app.services.milestones import jst_today, milestone_at
    from tests.conftest import TestingSessionLocal

    today = jst_today()
    key = today.strftime("%Y-%m-%d")
    with TestingSessionLocal() as s:
        task = Task(project_id=1, wbs_code="ms-rel", name="関連工程",
                    planned_start_at=milestone_at(key),
                    planned_finish_at=milestone_at((today + timedelta(days=3)).strftime("%Y-%m-%d")))
        s.add(task)
        s.flush()
        d, _ = _ms_fixture(s, order=220, type_name="日単位", planned_at=milestone_at(key, "day"),
                           schedule_precision="day", related_task_id=task.id)
        am, _ = _ms_fixture(s, order=221, type_name="午前", planned_at=milestone_at(key, "half_day", "AM"),
                            schedule_precision="half_day")
        pm, _ = _ms_fixture(s, order=222, type_name="午後", planned_at=milestone_at(key, "half_day", "PM"),
                            schedule_precision="half_day")
        s.commit()
        ids = (d.id, am.id, pm.id, task.id)

    data = _ms_get(client, "admin@test.jp", "include_candidates=false&limit=5000")
    by_id = {m["id"]: m for m in data["milestones"]}

    assert by_id[ids[0]]["schedule_precision"] == "day"
    assert _jst_wall(by_id[ids[0]]["planned_at"]).hour == 0
    assert by_id[ids[1]]["schedule_precision"] == "half_day"
    assert _jst_wall(by_id[ids[1]]["planned_at"]).hour == 0
    assert by_id[ids[2]]["schedule_precision"] == "half_day"
    assert _jst_wall(by_id[ids[2]]["planned_at"]).hour == 12, "午後は12:00で返ること"
    # 9時間ずれ（日付が前後の日へ動いていない）
    for mid in ids[:3]:
        assert _jst_wall(by_id[mid]["planned_at"]).strftime("%Y-%m-%d") == key

    # 関連工程は同じ案件のもの
    rel = by_id[ids[0]]
    assert rel["related_task_id"] == ids[3]
    assert rel["project_id"] == 1
    # 関連付けが無いものは工程との関係を推測しない
    assert by_id[ids[1]]["related_task_id"] is None
    assert by_id[ids[1]]["related_task_conflict"] is False


def test_milestone_api_has_no_n_plus_one(client):
    """件数を増やしてもSQL発行回数が比例増加しないこと。"""
    from sqlalchemy import event

    from app.services.milestones import jst_today, milestone_at
    from tests.conftest import TestingSessionLocal, engine

    today = jst_today()

    def count_queries(qs: str) -> int:
        statements: list[str] = []

        def before(conn, cursor, statement, params, context, executemany):
            statements.append(statement)

        event.listen(engine, "before_cursor_execute", before)
        try:
            _ms_get(client, "admin@test.jp", qs)
        finally:
            event.remove(engine, "before_cursor_execute", before)
        return len(statements)

    qs = "include_candidates=false&limit=5000"
    before_rows = _ms_get(client, "admin@test.jp", qs)["registered_count"]
    baseline = count_queries(qs)

    # 30件追加しても発行回数は変わらないはず
    with TestingSessionLocal() as s:
        for i in range(30):
            _ms_fixture(s, order=300 + i, type_name=f"N+1確認{i}",
                        planned_at=milestone_at(today.strftime("%Y-%m-%d")))
        s.commit()

    after_rows = _ms_get(client, "admin@test.jp", qs)["registered_count"]
    grown = count_queries(qs)

    assert after_rows >= before_rows + 30, "件数が増えていること"
    assert grown == baseline, f"件数を増やすとSQL発行回数が増えた: {baseline} → {grown}"


def test_milestone_api_applies_all_five_roles(client):
    """5権限それぞれで案件スコープが正しく効くこと。"""
    from app.core.security import hash_password
    from app.models import ProjectMember, User
    from tests.conftest import TestingSessionLocal

    with TestingSessionLocal() as s:
        for email, role in [("qm@test.jp", "QUALITY_MANAGER"), ("viewer@test.jp", "VIEWER"),
                            ("fw@test.jp", "FIELD_WORKER")]:
            if s.query(User).filter_by(email=email).one_or_none() is None:
                s.add(User(email=email, hashed_password=hash_password("pass"), name=role, role=role))
        s.flush()
        # FIELD_WORKER は案件2だけに割り当てる
        fw = s.query(User).filter_by(email="fw@test.jp").one()
        if s.query(ProjectMember).filter_by(user_id=fw.id).one_or_none() is None:
            s.add(ProjectMember(project_id=2, user_id=fw.id, role="FIELD_WORKER"))
        s.commit()

    # 内部ロールは全案件を見られる
    for email in ("admin@test.jp", "pm@test.jp", "qm@test.jp", "viewer@test.jp"):
        data = _ms_get(client, email, "include_candidates=false&limit=5000")
        assert {m["project_id"] for m in data["milestones"]} == {1, 2}, f"{email} が全案件を見られない"

    # 協力会社ロールは割当案件だけ
    partner = _ms_get(client, "partner@test.jp", "include_candidates=false&limit=5000")
    assert {m["project_id"] for m in partner["milestones"]} == {1}
    fw_data = _ms_get(client, "fw@test.jp", "include_candidates=false&limit=5000")
    assert {m["project_id"] for m in fw_data["milestones"]} == {2}

    # 件数・選択肢・候補のすべてに同じスコープが効く
    t = token(client, "fw@test.jp")
    opts = client.get("/api/schedule/milestones/options", headers={"Authorization": f"Bearer {t}"}).json()
    assert [p["construction_number"] for p in opts["projects"]] == ["T-002"]
    with_cand = _ms_get(client, "fw@test.jp", "limit=5000")
    assert all(c["project_id"] == 2 for c in with_cand["candidates"])
    summary = client.get("/api/schedule/milestones/summary", headers={"Authorization": f"Bearer {t}"}).json()
    assert summary["registered_count"] == fw_data["registered_count"]


def test_milestone_api_conflict_with_related_task(client):
    """関連工程との日程矛盾を確定計算で検知し、絞り込みできること。"""
    from app.models import Task
    from app.services.milestones import jst_today, milestone_at
    from tests.conftest import TestingSessionLocal

    today = jst_today()
    with TestingSessionLocal() as s:
        task = Task(project_id=1, wbs_code="ms-conf", name="矛盾確認の工程",
                    planned_start_at=milestone_at(today.strftime("%Y-%m-%d")),
                    planned_finish_at=milestone_at((today + timedelta(days=3)).strftime("%Y-%m-%d")))
        s.add(task)
        s.flush()
        inside, _ = _ms_fixture(s, order=400, type_name="期間内",
                                planned_at=milestone_at((today + timedelta(days=1)).strftime("%Y-%m-%d")),
                                related_task_id=task.id)
        outside, _ = _ms_fixture(s, order=401, type_name="期間外",
                                 planned_at=milestone_at((today + timedelta(days=10)).strftime("%Y-%m-%d")),
                                 related_task_id=task.id)
        s.commit()
        inside_id, outside_id = inside.id, outside.id

    data = _ms_get(client, "admin@test.jp", "include_candidates=false&limit=5000")
    by_id = {m["id"]: m for m in data["milestones"]}
    assert by_id[inside_id]["related_task_conflict"] is False
    assert by_id[outside_id]["related_task_conflict"] is True

    only = _ms_get(client, "admin@test.jp", "conflict_only=true&include_candidates=false&limit=5000")
    ids = {m["id"] for m in only["milestones"]}
    assert outside_id in ids and inside_id not in ids
    assert all(m["related_task_conflict"] for m in only["milestones"])


# --- Ver.0.3 Phase 3 (P3-2補足): CRUD と Excel/PDF 出力 ---------------------
def _auth(client, email):
    return {"Authorization": f"Bearer {token(client, email)}"}


def test_milestone_crud_lifecycle_and_audit(client):
    """登録→取得→更新→再取得→論理削除と、監査ログの記録。"""
    from app.models import AuditLog, Milestone, MilestoneType
    from tests.conftest import TestingSessionLocal

    h = _auth(client, "pm@test.jp")
    with TestingSessionLocal() as s:
        mt = MilestoneType(code="crud-種別", name="CRUD種別", sort_order=500)
        s.add(mt)
        s.commit()
        type_id = mt.id

    created = client.post("/api/schedule/milestones", headers=h, json={
        "project_id": 1, "milestone_type_id": type_id, "name": "CRUD対象",
        "planned_at": "2026-09-10T00:00:00+09:00", "schedule_precision": "day", "status": "予定",
    })
    assert created.status_code == 201, created.text
    mid = created.json()["id"]
    assert created.json()["record_kind"] == "milestone"

    got = client.get(f"/api/schedule/milestones/{mid}", headers=h)
    assert got.status_code == 200 and got.json()["name"] == "CRUD対象"

    updated = client.put(f"/api/schedule/milestones/{mid}", headers=h, json={
        "name": "CRUD対象（更新後）", "status": "完了",
        "actual_at": "2026-09-12T00:00:00+09:00", "change_reason": "テスト更新",
    })
    assert updated.status_code == 200, updated.text
    body = updated.json()
    assert body["name"] == "CRUD対象（更新後）"
    assert body["is_completed"] and body["was_delayed"] and body["delay_days"] == 2

    again = client.get(f"/api/schedule/milestones/{mid}", headers=h).json()
    assert again["name"] == "CRUD対象（更新後）" and again["status"] == "完了"

    assert client.delete(f"/api/schedule/milestones/{mid}?reason=テスト削除", headers=h).status_code == 204
    assert client.get(f"/api/schedule/milestones/{mid}", headers=h).status_code == 404

    with TestingSessionLocal() as s:
        row = s.get(Milestone, mid)
        assert row is not None, "物理削除ではなく論理削除であること"
        assert row.deleted_at is not None and row.delete_reason == "テスト削除"
        actions = {a.action for a in s.query(AuditLog).filter_by(entity_type="milestone",
                                                                 entity_id=str(mid)).all()}
        assert {"CREATE", "UPDATE", "DELETE"} <= actions, f"監査ログが足りない: {actions}"

    # 一覧からも消える
    assert all(m["id"] != mid for m in _ms_get(client, "pm@test.jp", "include_candidates=false&limit=5000")["milestones"])


def test_milestone_create_from_candidate_removes_it(client):
    """未設定候補から登録すると、その候補が消えること。"""
    from app.models import MilestoneType
    from tests.conftest import TestingSessionLocal

    with TestingSessionLocal() as s:
        mt = MilestoneType(code="cand-登録元", name="候補から登録", sort_order=510)
        s.add(mt)
        s.commit()
        type_id = mt.id

    before = _ms_get(client, "pm@test.jp", "project_ids=1&limit=5000")
    assert any(c["milestone_type_id"] == type_id for c in before["candidates"])

    h = _auth(client, "pm@test.jp")
    r = client.post("/api/schedule/milestones", headers=h, json={
        "project_id": 1, "milestone_type_id": type_id, "name": "候補から作成",
        "planned_at": "2026-10-01T00:00:00+09:00",
    })
    assert r.status_code == 201, r.text
    # 候補は実レコードではないため、予定日・状態・担当者は入力値だけが入る
    assert r.json()["responsible_id"] is None and r.json()["company_id"] is None

    after = _ms_get(client, "pm@test.jp", "project_ids=1&limit=5000")
    assert all(c["milestone_type_id"] != type_id for c in after["candidates"])
    assert after["candidate_count"] == before["candidate_count"] - 1
    assert after["registered_count"] == before["registered_count"] + 1


def test_milestone_crud_reference_validation(client):
    """存在しない/削除済み/別案件は422、スコープ外は403で拒否すること。"""
    from datetime import datetime, timezone as _tz

    from app.models import MilestoneType, Task
    from tests.conftest import TestingSessionLocal

    h = _auth(client, "pm@test.jp")
    with TestingSessionLocal() as s:
        mt = MilestoneType(code="ref-種別", name="参照検証", sort_order=520)
        inactive = MilestoneType(code="ref-非active", name="無効な区分", sort_order=521, active=False)
        t_other = Task(project_id=2, wbs_code="ref-2", name="別案件の工程")
        t_gone = Task(project_id=1, wbs_code="ref-del", name="削除済み工程")
        s.add_all([mt, inactive, t_other, t_gone])
        s.flush()
        t_gone.deleted_at = datetime.now(_tz.utc)
        s.commit()
        type_id, inactive_id, other_task, gone_task = mt.id, inactive.id, t_other.id, t_gone.id

    base = {"project_id": 1, "milestone_type_id": type_id, "name": "参照検証",
            "planned_at": "2026-09-10T00:00:00+09:00"}

    # 存在しない参照 → 422
    for field, value in [("milestone_type_id", 999999), ("responsible_id", 999999),
                         ("company_id", 999999), ("related_task_id", 999999)]:
        r = client.post("/api/schedule/milestones", headers=h, json={**base, field: value})
        assert r.status_code == 422, f"{field}: {r.status_code}"
    # 別案件の工程 → 422 / 削除済みの工程 → 422
    assert client.post("/api/schedule/milestones", headers=h,
                       json={**base, "related_task_id": other_task}).status_code == 422
    assert client.post("/api/schedule/milestones", headers=h,
                       json={**base, "related_task_id": gone_task}).status_code == 422
    # 非active の区分は新規登録に使えない
    assert client.post("/api/schedule/milestones", headers=h,
                       json={**base, "milestone_type_id": inactive_id}).status_code == 422
    # 存在しない案件 → 422
    assert client.post("/api/schedule/milestones", headers=h,
                       json={**base, "project_id": 999999}).status_code == 422
    # スコープ外の案件 → 403（協力会社ユーザーは案件1のみ）
    hp = _auth(client, "partner@test.jp")
    assert client.post("/api/schedule/milestones", headers=hp,
                       json={**base, "project_id": 2}).status_code in (403, 404)

    # 更新でも同じ検証が効く
    ok = client.post("/api/schedule/milestones", headers=h, json=base)
    assert ok.status_code == 201
    mid = ok.json()["id"]
    assert client.put(f"/api/schedule/milestones/{mid}", headers=h,
                      json={"related_task_id": other_task}).status_code == 422
    assert client.put(f"/api/schedule/milestones/{mid}", headers=h,
                      json={"related_task_id": gone_task}).status_code == 422
    # 既存レコードは非activeの区分でも更新できる（履歴を壊さない）
    with TestingSessionLocal() as s:
        s.get(__import__("app.models", fromlist=["Milestone"]).Milestone, mid).milestone_type_id = inactive_id
        s.commit()
    assert client.get(f"/api/schedule/milestones/{mid}", headers=h).status_code == 200
    assert client.put(f"/api/schedule/milestones/{mid}", headers=h,
                      json={"name": "非active区分のまま更新"}).status_code == 200


def test_milestone_update_cannot_change_project(client):
    """更新で案件を無断で変更できないこと。"""
    from app.models import MilestoneType
    from tests.conftest import TestingSessionLocal

    h = _auth(client, "pm@test.jp")
    with TestingSessionLocal() as s:
        mt = MilestoneType(code="proj-固定", name="案件固定", sort_order=530)
        s.add(mt)
        s.commit()
        type_id = mt.id

    mid = client.post("/api/schedule/milestones", headers=h, json={
        "project_id": 1, "milestone_type_id": type_id, "name": "案件は変えない",
        "planned_at": "2026-09-10T00:00:00+09:00"}).json()["id"]

    # project_id は MilestoneUpdate に無いため、送っても無視される
    r = client.put(f"/api/schedule/milestones/{mid}", headers=h, json={"project_id": 2, "name": "変更後"})
    assert r.status_code == 200
    assert r.json()["project_id"] == 1, "案件が変わってしまった"


def test_milestone_crud_precision_roundtrip_via_api(client):
    """API経由で day / half_day 午前・午後が往復し、9時間ずれないこと。"""
    from app.models import MilestoneType
    from tests.conftest import TestingSessionLocal

    h = _auth(client, "pm@test.jp")
    with TestingSessionLocal() as s:
        mt = MilestoneType(code="prec-api", name="粒度API", sort_order=540)
        s.add(mt)
        s.commit()
        type_id = mt.id

    cases = [
        ("day", "2026-09-10T00:00:00+09:00", 0, "2026-09-10"),
        ("half_day", "2026-09-11T00:00:00+09:00", 0, "2026-09-11"),
        ("half_day", "2026-09-12T12:00:00+09:00", 12, "2026-09-12"),
    ]
    for precision, planned, hour, day in cases:
        r = client.post("/api/schedule/milestones", headers=h, json={
            "project_id": 1, "milestone_type_id": type_id, "name": f"{precision}-{hour}",
            "planned_at": planned, "schedule_precision": precision})
        assert r.status_code == 201, r.text
        got = client.get(f"/api/schedule/milestones/{r.json()['id']}", headers=h).json()
        assert got["schedule_precision"] == precision
        w = _jst_wall(got["planned_at"])
        assert w.hour == hour, f"{precision}/{hour}: {got['planned_at']}"
        assert w.strftime("%Y-%m-%d") == day, f"9時間ずれ: {got['planned_at']}"


def test_milestone_roles_can_write(client):
    """作成・更新・削除の可否がロールで分かれること。"""
    from app.core.security import hash_password
    from app.models import MilestoneType, User
    from tests.conftest import TestingSessionLocal

    with TestingSessionLocal() as s:
        mt = MilestoneType(code="role-種別", name="権限確認", sort_order=550)
        s.add(mt)
        for email, role in [("qm2@test.jp", "QUALITY_MANAGER"), ("viewer2@test.jp", "VIEWER")]:
            if s.query(User).filter_by(email=email).one_or_none() is None:
                s.add(User(email=email, hashed_password=hash_password("pass"), name=role, role=role))
        s.commit()
        type_id = mt.id

    body = {"project_id": 1, "milestone_type_id": type_id, "name": "権限確認",
            "planned_at": "2026-09-10T00:00:00+09:00"}

    # 作成できる（ADMIN は require_roles を常に通る）
    for email in ("admin@test.jp", "pm@test.jp"):
        r = client.post("/api/schedule/milestones", headers=_auth(client, email), json=body)
        assert r.status_code == 201, f"{email}: {r.text}"
        mid = r.json()["id"]
        assert client.put(f"/api/schedule/milestones/{mid}", headers=_auth(client, email),
                          json={"name": "更新"}).status_code == 200
        assert client.delete(f"/api/schedule/milestones/{mid}", headers=_auth(client, email)).status_code == 204

    # 参照専用のロールは書き込めない
    for email in ("qm2@test.jp", "viewer2@test.jp", "partner@test.jp"):
        assert client.post("/api/schedule/milestones", headers=_auth(client, email),
                           json=body).status_code == 403, email


def test_milestone_export_matches_screen(client):
    """Excel/PDF が画面と同じ条件・同じ件数で生成され、スコープが効くこと。"""
    h = _auth(client, "admin@test.jp")
    qs = "actual=missing&include_candidates=true&limit=5000"
    screen = _ms_get(client, "admin@test.jp", qs)

    for fmt in ("xlsx", "pdf"):
        r = client.get(f"/api/schedule/milestones/export?format={fmt}&{qs}", headers=h)
        assert r.status_code == 200, r.text
        assert len(r.content) > 500
        assert "attachment;" in r.headers["content-disposition"]

    # 出力は collect() を通すので、同じ条件なら件数が一致する
    from app.api.milestones import MilestoneFilters, collect
    from app.models import User as U
    from tests.conftest import TestingSessionLocal

    with TestingSessionLocal() as s:
        admin = s.query(U).filter_by(email="admin@test.jp").one()
        partner = s.query(U).filter_by(email="partner@test.jp").one()
        f = MilestoneFilters(actual="missing", include_candidates=True, limit=5000)
        same = collect(s, admin, f)
        assert same.registered_count == screen["registered_count"]
        assert same.candidate_count == screen["candidate_count"]
        assert [m.id for m in same.milestones] == [m["id"] for m in screen["milestones"]], "並び順も一致"

        # 出力にも案件スコープが効く
        scoped = collect(s, partner, MilestoneFilters(include_candidates=True, limit=5000))
        assert {m.project_id for m in scoped.milestones} == {1}
        assert all(c.project_id == 1 for c in scoped.candidates)


def test_milestone_multiple_records_same_project_and_type(client):
    """同一案件・同一区分が複数件あっても1件に潰れないこと。"""
    from app.models import MilestoneType
    from tests.conftest import TestingSessionLocal

    h = _auth(client, "pm@test.jp")
    with TestingSessionLocal() as s:
        mt = MilestoneType(code="dup-種別", name="重複可能な区分", sort_order=560)
        s.add(mt)
        s.commit()
        type_id = mt.id

    ids = []
    for i, day in enumerate(("2026-09-20", "2026-09-21")):
        r = client.post("/api/schedule/milestones", headers=h, json={
            "project_id": 1, "milestone_type_id": type_id, "name": f"同一区分{i + 1}",
            "planned_at": f"{day}T00:00:00+09:00"})
        assert r.status_code == 201, r.text
        ids.append(r.json()["id"])

    data = _ms_get(client, "pm@test.jp", f"milestone_type_ids={type_id}&project_ids=1&limit=5000")
    got = [m for m in data["milestones"] if m["id"] in ids]
    assert len(got) == 2, "同一案件・同一区分の複数件が欠落した"
    assert data["registered_count"] >= 2
    # 登録済みがあるので、その区分は候補に出ない
    assert all(c["milestone_type_id"] != type_id for c in data["candidates"])


# --- Ver.0.3 Phase 3 (P3-4): 表示名依存の判定を milestones へ置き換え ---------
def test_task_named_delivery_is_not_a_milestone(client):
    """「引き渡し」という名前の通常工程は、マイルストーンとして扱われないこと。

    工程名の文字列一致（name == '引き渡し'）をやめ、milestones を正データにした
    ことの確認。工程はあくまで tasks 側にだけ存在する。
    """
    from datetime import datetime, timezone as _tz

    from app.models import Task
    from tests.conftest import TestingSessionLocal

    h = _auth(client, "pm@test.jp")
    with TestingSessionLocal() as s:
        task = Task(project_id=1, wbs_code="P34.1", name="引き渡し", status="未着手",
                    planned_start_at=datetime(2026, 9, 1, tzinfo=_tz.utc),
                    planned_finish_at=datetime(2026, 9, 2, tzinfo=_tz.utc))
        s.add(task)
        s.commit()
        task_id = task.id

    data = _ms_get(client, "pm@test.jp", "project_ids=1&limit=5000")
    # 工程はマイルストーン一覧に現れない（IDが衝突しても別物として扱われる）
    assert all(m["record_kind"] == "milestone" for m in data["milestones"])
    assert not any(m["name"] == "引き渡し" and m["related_task_id"] is None
                   and m["milestone_type_id"] is None for m in data["milestones"])

    # 工程一覧側では通常の工程として返る（バー表示に必要な期間を持つ）
    r = client.get("/api/projects/1/tasks", headers=h)
    assert r.status_code == 200
    hit = [t for t in r.json() if t["id"] == task_id]
    assert len(hit) == 1
    assert hit[0]["planned_start_at"] and hit[0]["planned_finish_at"]


def test_task_named_delivery_can_be_moved(client):
    """「引き渡し」という名前でも、工程の日程変更が名前を理由に拒否されないこと。"""
    from datetime import datetime, timezone as _tz

    from app.models import Task
    from app.services.milestones import as_jst
    from tests.conftest import TestingSessionLocal

    h = _auth(client, "pm@test.jp")
    with TestingSessionLocal() as s:
        task = Task(project_id=1, wbs_code="P34.2", name="引き渡し", status="未着手",
                    planned_start_at=datetime(2026, 9, 10, tzinfo=_tz.utc),
                    planned_finish_at=datetime(2026, 9, 11, tzinfo=_tz.utc))
        s.add(task)
        s.commit()
        task_id = task.id

    r = client.put(f"/api/tasks/{task_id}", headers=h, json={
        "name": "引き渡し",
        "planned_start_at": "2026-09-12T00:00:00+09:00",
        "planned_finish_at": "2026-09-13T00:00:00+09:00",
        "change_reason": "P3-4 ドラッグ相当の日程変更",
    })
    assert r.status_code == 200, r.text
    # 保存後の値を JST へ直して確認する（テストDBは naive、本番は timezone 付き）
    moved = as_jst(datetime.fromisoformat(r.json()["planned_start_at"]))
    assert (moved.year, moved.month, moved.day, moved.hour) == (2026, 9, 12, 0)


def test_milestone_kind_survives_rename(client):
    """マイルストーン名を変えても、種別と record_kind が変わらないこと。"""
    from app.models import MilestoneType
    from tests.conftest import TestingSessionLocal

    h = _auth(client, "pm@test.jp")
    with TestingSessionLocal() as s:
        mt = MilestoneType(code="p34-rename", name="立会検査", sort_order=610)
        s.add(mt)
        s.commit()
        type_id = mt.id

    created = client.post("/api/schedule/milestones", headers=h, json={
        "project_id": 1, "milestone_type_id": type_id, "name": "立会検査（初回）",
        "planned_at": "2026-10-01T00:00:00+09:00"})
    assert created.status_code == 201, created.text
    ms_id = created.json()["id"]

    renamed = client.put(f"/api/schedule/milestones/{ms_id}", headers=h,
                         json={"name": "まったく違う名前"})
    assert renamed.status_code == 200, renamed.text
    body = renamed.json()
    assert body["record_kind"] == "milestone"
    assert body["milestone_type_id"] == type_id
    assert body["milestone_type"] == "立会検査"

    data = _ms_get(client, "pm@test.jp", f"milestone_type_ids={type_id}&limit=5000")
    got = [m for m in data["milestones"] if m["id"] == ms_id]
    assert len(got) == 1 and got[0]["milestone_type_id"] == type_id


def test_milestone_candidates_are_never_returned_as_records(client):
    """未設定候補が実マイルストーンとして返らないこと（描画の取り違えを防ぐ）。"""
    data = _ms_get(client, "pm@test.jp", "limit=5000")
    assert all(m["record_kind"] == "milestone" for m in data["milestones"])
    assert all(c["record_kind"] == "candidate" for c in data["candidates"])
    # 候補は実在IDも日付も持たない
    for c in data["candidates"]:
        assert "id" not in c and "planned_at" not in c and "actual_at" not in c
    # 件数は登録済みと候補で別々に数える
    assert data["registered_count"] == len(data["milestones"])
    assert data["candidate_count"] == len(data["candidates"])


def test_milestone_soft_deleted_is_not_drawn(client):
    """論理削除したマイルストーンが一覧に出ないこと。"""
    from app.models import MilestoneType
    from tests.conftest import TestingSessionLocal

    h = _auth(client, "pm@test.jp")
    with TestingSessionLocal() as s:
        mt = MilestoneType(code="p34-deleted", name="削除確認", sort_order=620)
        s.add(mt)
        s.commit()
        type_id = mt.id

    created = client.post("/api/schedule/milestones", headers=h, json={
        "project_id": 1, "milestone_type_id": type_id, "name": "消える予定",
        "planned_at": "2026-10-05T00:00:00+09:00"})
    ms_id = created.json()["id"]
    assert any(m["id"] == ms_id for m in _ms_get(client, "pm@test.jp", "limit=5000")["milestones"])

    assert client.delete(f"/api/schedule/milestones/{ms_id}", headers=h).status_code == 204
    after = _ms_get(client, "pm@test.jp", "limit=5000")
    assert all(m["id"] != ms_id for m in after["milestones"])


def test_milestone_inactive_type_is_kept_as_history(client):
    """非activeの区分を参照する既存レコードも、履歴として区分名つきで返ること。"""
    from app.models import MilestoneType
    from tests.conftest import TestingSessionLocal

    h = _auth(client, "pm@test.jp")
    with TestingSessionLocal() as s:
        mt = MilestoneType(code="p34-inactive", name="廃止された区分", sort_order=630)
        s.add(mt)
        s.commit()
        type_id = mt.id

    created = client.post("/api/schedule/milestones", headers=h, json={
        "project_id": 1, "milestone_type_id": type_id, "name": "廃止区分の履歴",
        "planned_at": "2026-10-08T00:00:00+09:00"})
    assert created.status_code == 201, created.text
    ms_id = created.json()["id"]

    with TestingSessionLocal() as s:
        s.query(MilestoneType).filter_by(id=type_id).one().active = False
        s.commit()

    data = _ms_get(client, "pm@test.jp", "limit=5000")
    hit = [m for m in data["milestones"] if m["id"] == ms_id]
    assert len(hit) == 1, "非active区分の既存レコードが消えた"
    assert hit[0]["milestone_type_id"] == type_id
    assert hit[0]["milestone_type"] == "廃止された区分"
    # 非active区分は新規登録の候補には出さない
    assert all(c["milestone_type_id"] != type_id for c in data["candidates"])


def test_milestone_same_name_different_id_are_separated(client):
    """同名でIDが異なるマイルストーンを取り違えないこと。"""
    from app.models import MilestoneType
    from tests.conftest import TestingSessionLocal

    h = _auth(client, "pm@test.jp")
    with TestingSessionLocal() as s:
        for code, order in (("p34-same-a", 640), ("p34-same-b", 641)):
            if s.query(MilestoneType).filter_by(code=code).one_or_none() is None:
                s.add(MilestoneType(code=code, name="同じ区分名", sort_order=order))
        s.commit()
        ids = [t.id for t in s.query(MilestoneType)
               .filter(MilestoneType.code.in_(["p34-same-a", "p34-same-b"])).all()]

    made = []
    for type_id in ids:
        r = client.post("/api/schedule/milestones", headers=h, json={
            "project_id": 1, "milestone_type_id": type_id, "name": "同じ名前",
            "planned_at": "2026-10-12T00:00:00+09:00"})
        assert r.status_code == 201, r.text
        made.append((r.json()["id"], type_id))

    data = _ms_get(client, "pm@test.jp", "limit=5000")
    by_id = {m["id"]: m for m in data["milestones"]}
    for ms_id, type_id in made:
        assert by_id[ms_id]["milestone_type_id"] == type_id, "同名の区分がIDで分離されていない"
    assert len({i for i, _ in made}) == 2


def test_milestone_plan_and_actual_are_separate_values(client):
    """予定日と実績日が同日でも、それぞれ別の値として返ること。"""
    from datetime import datetime

    from app.models import MilestoneType
    from app.services.milestones import as_jst
    from tests.conftest import TestingSessionLocal

    h = _auth(client, "pm@test.jp")
    with TestingSessionLocal() as s:
        mt = MilestoneType(code="p34-both", name="予実同日", sort_order=650)
        s.add(mt)
        s.commit()
        type_id = mt.id

    r = client.post("/api/schedule/milestones", headers=h, json={
        "project_id": 1, "milestone_type_id": type_id, "name": "同日の予実",
        "schedule_precision": "half_day",
        "planned_at": "2026-10-15T00:00:00+09:00",   # 午前
        "actual_at": "2026-10-15T12:00:00+09:00"})   # 午後
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["planned_at"] != body["actual_at"]
    assert as_jst(datetime.fromisoformat(body["planned_at"])).hour == 0
    assert as_jst(datetime.fromisoformat(body["actual_at"])).hour == 12
    assert body["is_completed"] is True


def test_project_schedule_scope_excludes_other_projects(client):
    """案件工程で使う project_id 指定に、別案件のマイルストーンが混ざらないこと。"""
    from app.models import MilestoneType
    from tests.conftest import TestingSessionLocal

    h = _auth(client, "pm@test.jp")
    with TestingSessionLocal() as s:
        mt = MilestoneType(code="p34-scope", name="スコープ確認", sort_order=660)
        s.add(mt)
        s.commit()
        type_id = mt.id
    for pid in (1, 2):
        assert client.post("/api/schedule/milestones", headers=h, json={
            "project_id": pid, "milestone_type_id": type_id, "name": f"案件{pid}の重要日",
            "planned_at": "2026-10-20T00:00:00+09:00"}).status_code == 201

    only1 = _ms_get(client, "pm@test.jp", "project_id=1&limit=5000")
    assert {m["project_id"] for m in only1["milestones"]} == {1}
    assert all(c["project_id"] == 1 for c in only1["candidates"])


def test_milestone_query_count_does_not_grow_with_rows(client):
    """行数が増えてもSQL発行回数が増えないこと（行ごと取得をしていない）。"""
    from sqlalchemy import event

    from app.models import MilestoneType
    from tests.conftest import TestingSessionLocal, engine

    h = _auth(client, "pm@test.jp")
    with TestingSessionLocal() as s:
        mt = MilestoneType(code="p34-n1", name="N+1確認", sort_order=670)
        s.add(mt)
        s.commit()
        type_id = mt.id

    def count_queries(qs: str) -> int:
        seen = []
        listener = lambda *a, **k: seen.append(1)  # noqa: E731
        event.listen(engine, "before_cursor_execute", listener)
        try:
            _ms_get(client, "pm@test.jp", qs)
        finally:
            event.remove(engine, "before_cursor_execute", listener)
        return len(seen)

    base = count_queries(f"milestone_type_ids={type_id}&limit=5000")
    for i in range(12):
        assert client.post("/api/schedule/milestones", headers=h, json={
            "project_id": 1, "milestone_type_id": type_id, "name": f"N+1確認{i}",
            "planned_at": "2026-10-25T00:00:00+09:00"}).status_code == 201
    after = count_queries(f"milestone_type_ids={type_id}&limit=5000")
    assert after == base, f"件数に比例してSQLが増えた（{base} → {after}）"


# --- Ver.0.3 Phase 3 (P3-5): カレンダーの共通イベントAPI ---------------------
def _cal(client, email, qs=""):
    t = token(client, email)
    r = client.get(f"/api/schedule/calendar/events?{qs}", headers={"Authorization": f"Bearer {t}"})
    assert r.status_code == 200, r.text
    return r.json()


def _cal_fixture(session):
    """5種別すべての元データを1日ぶん用意する（カレンダー専用テーブルは作らない）。"""
    from datetime import date as _date, datetime as _dt

    from app.models import DailyReport, Milestone, MilestoneType, QualityCheck, Task, TestRecord
    from app.services.milestones import JST, milestone_at

    mt = session.query(MilestoneType).filter_by(code="p35-cal").one_or_none()
    if mt is None:
        mt = MilestoneType(code="p35-cal", name="カレンダー確認", sort_order=700)
        session.add(mt)
        session.flush()
    # 日時はすべて JST で組み立てる（Phase 1 と同じ意味論）
    task = Task(project_id=1, wbs_code="P35.1", name="カレンダー工程", status="施工中",
                planned_start_at=_dt(2026, 11, 10, tzinfo=JST),
                planned_finish_at=_dt(2026, 11, 12, tzinfo=JST),   # exclusive（11/11まで）
                actual_start_at=_dt(2026, 11, 10, tzinfo=JST))
    ms = Milestone(project_id=1, milestone_type_id=mt.id, name="カレンダー重要日",
                   planned_at=milestone_at("2026-11-10", "half_day", "PM"),  # JST 12:00
                   status="予定", schedule_precision="half_day")
    qc = QualityCheck(project_id=1, inspect_item="カレンダー品質確認", status="確認待ち",
                      due_date=_date(2026, 11, 10))
    dr = DailyReport(project_id=1, report_date=_date(2026, 11, 10), status="DRAFT",
                     place="カレンダー現場")
    tr = TestRecord(project_id=1, test_type="カレンダー試験", judge="未判定",
                    measured_at=_dt(2026, 11, 10, 10, 30, tzinfo=JST))
    session.add_all([task, ms, qc, dr, tr])
    session.commit()
    return {"task": task.id, "milestone": ms.id, "quality_check": qc.id,
            "daily_report": dr.id, "test_record": tr.id, "type_id": mt.id}


def test_calendar_integrates_all_real_sources(client):
    """実在する5種別だけを共通形式へ変換して返すこと。"""
    from tests.conftest import TestingSessionLocal

    with TestingSessionLocal() as s:
        ids = _cal_fixture(s)

    data = _cal(client, "pm@test.jp", "date_from=2026-11-10&date_to=2026-11-10&project_ids=1")
    kinds = {(e["source_kind"], e["record_kind"]) for e in data["events"]}
    assert ("task", "plan") in kinds and ("task", "actual") in kinds
    assert ("milestone", "plan") in kinds
    assert ("quality_check", "due") in kinds
    assert ("daily_report", "actual") in kinds
    assert ("test_record", "actual") in kinds

    by_kind = {e["source_kind"]: e for e in data["events"]}
    # 元レコードへ戻れること
    assert by_kind["task"]["source_url"] == "/projects/1/schedule"
    assert by_kind["milestone"]["source_url"].startswith("/projects/1/schedule/milestones?types=")
    assert by_kind["quality_check"]["source_url"] == "/quality"
    assert by_kind["daily_report"]["source_url"] == "/daily-report"
    # 元IDを保持している
    assert by_kind["quality_check"]["source_id"] == ids["quality_check"]
    assert data["total"] == data["displayed"] and data["truncated"] is False


def test_calendar_event_id_never_collides_across_tables(client):
    """別テーブルの同じ数値IDを1件に取り違えないこと。"""
    data = _cal(client, "pm@test.jp", "date_from=2026-01-01&date_to=2026-12-31&limit=5000")
    assert len({e["event_id"] for e in data["events"]}) == len(data["events"])
    for e in data["events"]:
        assert e["event_id"] == f"{e['source_kind']}:{e['record_kind']}:{e['source_id']}"
    # 数値IDだけなら重複しうる組み合わせが、実際に別イベントとして残っている
    numeric = [(e["source_id"], e["source_kind"]) for e in data["events"]]
    dup_ids = {i for i, _ in numeric if sum(1 for j, _ in numeric if j == i) > 1}
    assert dup_ids, "同じ数値IDを持つ別種別のデータが無く、この確認が成立していない"


def test_calendar_jst_day_boundary(client):
    """JSTの日付境界でずれないこと（9時間ずれの検出）。"""
    from datetime import datetime

    from app.services.milestones import as_jst
    from tests.conftest import TestingSessionLocal

    with TestingSessionLocal() as s:
        ids = _cal_fixture(s)

    # JST 11/10 の1日だけを指定して、その日のイベントが取れる
    same = _cal(client, "pm@test.jp", "date_from=2026-11-10&date_to=2026-11-10&project_ids=1")
    assert any(e["source_id"] == ids["milestone"] and e["source_kind"] == "milestone"
               for e in same["events"])
    # 前日を指定すると入らない（UTC基準で判定していれば混入する）
    before = _cal(client, "pm@test.jp", "date_from=2026-11-09&date_to=2026-11-09&project_ids=1")
    assert not any(e["source_kind"] == "milestone" and e["source_id"] == ids["milestone"]
                   for e in before["events"])
    # half_day の午後は JST 12:00 のまま返る
    ms = [e for e in same["events"] if e["source_kind"] == "milestone"][0]
    assert as_jst(datetime.fromisoformat(ms["start_at"])).hour == 12
    assert ms["schedule_precision"] == "half_day"


def test_calendar_range_is_half_open(client):
    """期間は Phase 1 と同じ [開始, 終了)。終了日は「その日を含む」指定。"""
    from tests.conftest import TestingSessionLocal

    with TestingSessionLocal() as s:
        ids = _cal_fixture(s)

    # 工程の予定は JST 11/10 00:00 〜 11/12 00:00（＝11/11 まで）
    inside = _cal(client, "pm@test.jp", "date_from=2026-11-11&date_to=2026-11-11&project_ids=1&source_kinds=task")
    assert any(e["source_id"] == ids["task"] and e["record_kind"] == "plan" for e in inside["events"])
    outside = _cal(client, "pm@test.jp", "date_from=2026-11-12&date_to=2026-11-12&project_ids=1&source_kinds=task")
    assert not any(e["source_id"] == ids["task"] and e["record_kind"] == "plan"
                   for e in outside["events"]), "終了日時（exclusive）の当日を含めてしまっている"


def test_calendar_excludes_deleted_and_out_of_scope(client):
    """論理削除済み・案件スコープ外を返さないこと。"""
    from datetime import datetime, timezone as _tz

    from app.models import Milestone
    from tests.conftest import TestingSessionLocal

    with TestingSessionLocal() as s:
        ids = _cal_fixture(s)

    qs = "date_from=2026-11-10&date_to=2026-11-10"
    assert any(e["source_kind"] == "milestone" and e["source_id"] == ids["milestone"]
               for e in _cal(client, "pm@test.jp", qs)["events"])

    with TestingSessionLocal() as s:
        row = s.get(Milestone, ids["milestone"])
        row.deleted_at = datetime.now(_tz.utc)
        s.commit()
    assert not any(e["source_kind"] == "milestone" and e["source_id"] == ids["milestone"]
                   for e in _cal(client, "pm@test.jp", qs)["events"])

    # 協力会社（割当は案件1のみ）は案件2のイベントを見られない
    partner = _cal(client, "partner@test.jp", "date_from=2026-01-01&date_to=2026-12-31&limit=5000")
    assert {e["project_id"] for e in partner["events"]} <= {1}


def test_calendar_same_day_same_name_not_merged(client):
    """同日・同名でもIDが違えば別イベントとして残ること。"""
    from datetime import date as _date

    from app.models import QualityCheck
    from tests.conftest import TestingSessionLocal

    with TestingSessionLocal() as s:
        s.add_all([
            QualityCheck(project_id=1, inspect_item="同名の確認", status="確認待ち",
                         due_date=_date(2026, 11, 20)),
            QualityCheck(project_id=1, inspect_item="同名の確認", status="確認待ち",
                         due_date=_date(2026, 11, 20)),
        ])
        s.commit()

    data = _cal(client, "pm@test.jp", "date_from=2026-11-20&date_to=2026-11-20&project_ids=1")
    same = [e for e in data["events"] if e["title"] == "同名の確認"]
    assert len(same) == 2, "同名の別レコードが1件に潰れた"
    assert same[0]["event_id"] != same[1]["event_id"]


def test_calendar_filters_are_and(client):
    """複合フィルターがANDで効くこと。"""
    from tests.conftest import TestingSessionLocal

    with TestingSessionLocal() as s:
        _cal_fixture(s)

    base = "date_from=2026-11-01&date_to=2026-11-30&project_ids=1"
    only_task = _cal(client, "pm@test.jp", f"{base}&source_kinds=task")
    assert {e["source_kind"] for e in only_task["events"]} == {"task"}

    with_q = _cal(client, "pm@test.jp", f"{base}&source_kinds=task&q=カレンダー")
    assert all("カレンダー" in e["title"] for e in with_q["events"])
    assert len(with_q["events"]) <= len(only_task["events"])

    impossible = _cal(client, "pm@test.jp", f"{base}&source_kinds=task&statuses=存在しない状態")
    assert impossible["events"] == [] and impossible["total"] == 0


def test_calendar_truncates_with_limit(client):
    """上限を超えたら truncated を立て、total と displayed を別に返すこと。"""
    data = _cal(client, "pm@test.jp", "date_from=2026-01-01&date_to=2026-12-31&limit=3")
    assert data["displayed"] == 3
    assert data["total"] > 3
    assert data["truncated"] is True
    assert data["limit"] == 3


def test_calendar_query_count_does_not_grow_with_rows(client):
    """件数が増えてもSQL発行回数が増えないこと（種別ごとに1本）。"""
    from datetime import date as _date

    from sqlalchemy import event

    from app.models import QualityCheck
    from tests.conftest import TestingSessionLocal, engine

    def count_queries(qs: str) -> int:
        seen = []
        listener = lambda *a, **k: seen.append(1)  # noqa: E731
        event.listen(engine, "before_cursor_execute", listener)
        try:
            _cal(client, "pm@test.jp", qs)
        finally:
            event.remove(engine, "before_cursor_execute", listener)
        return len(seen)

    qs = "date_from=2026-12-01&date_to=2026-12-31&project_ids=1&limit=5000"
    base = count_queries(qs)
    with TestingSessionLocal() as s:
        for i in range(15):
            s.add(QualityCheck(project_id=1, inspect_item=f"N+1確認{i}", status="確認待ち",
                               due_date=_date(2026, 12, 10)))
        s.commit()
    after = count_queries(qs)
    assert after == base, f"件数に比例してSQLが増えた（{base} → {after}）"


def test_calendar_options_use_ids(client):
    """選択肢は同名でも区別できるよう id と name の組で返すこと。"""
    t = token(client, "pm@test.jp")
    r = client.get("/api/schedule/calendar/options", headers={"Authorization": f"Bearer {t}"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert {k["key"] for k in body["source_kinds"]} == {
        "task", "milestone", "quality_check", "daily_report", "test_record"}
    for key in ("projects", "responsibles", "companies"):
        assert all("id" in v and "name" in v for v in body[key])


def test_calendar_has_no_dedicated_table(client):
    """カレンダー専用テーブルへ二重保存していないこと。"""
    from app.core.db import Base

    names = set(Base.metadata.tables)
    assert not {n for n in names if "calendar" in n}, "カレンダー専用テーブルを作っている"


# ---------------------------------------------------------------------------
# Phase 4: 案件ライフサイクル（検索・登録・更新・権限）
#
# Seed の固定ID には依存せず、この節が自分で作った実データだけを対象にする。
# 案件は工事番号・工事名の接頭辞で識別し、他テストの案件と混ざらないようにする。
# ---------------------------------------------------------------------------

P4_TAG = "P4LC"  # 絞り込み確認用の案件群
P4_PERIOD_TAG = "P4PER"  # 期間検索の境界確認用の案件群

_p4_seq = [0]


def _p4_number(prefix: str = "P4NEW") -> str:
    """テスト内で一意な工事番号。Seed の固定値を使い回さない。"""
    _p4_seq[0] += 1
    return f"{prefix}-{_p4_seq[0]:04d}"


def _p4_head(client, email: str) -> dict:
    return {"Authorization": f"Bearer {token(client, email)}"}


def _p4_search(client, email: str, qs: str = "") -> dict:
    r = client.get(f"/api/projects/search?{qs}", headers=_p4_head(client, email))
    assert r.status_code == 200, r.text
    return r.json()


def _p4_ids(client, email: str, qs: str) -> list[int]:
    return [p["id"] for p in _p4_search(client, email, qs)["items"]]


def _p4_walk_pages(client, email: str, qs: str, per_page: int = 2):
    """全ページを順に取得し、(通し順のID列, total, ページ数) を返す。"""
    ids: list[int] = []
    totals: set[int] = set()
    page = 1
    pages = 1
    while True:
        data = _p4_search(client, email, f"{qs}&per_page={per_page}&page={page}")
        totals.add(data["total"])
        pages = data["pages"]
        ids += [p["id"] for p in data["items"]]
        if page >= pages:
            break
        page += 1
    assert len(totals) == 1, f"ページごとに total が変わっている: {totals}"
    return ids, totals.pop(), pages


@pytest.fixture(scope="module")
def p4():
    """Phase 4 のテストデータ。ダミーの固定IDは作らず、作成したIDを返す。"""
    from datetime import date

    from app.core.security import hash_password
    from app.models import Company, Department, Project, ProjectMember, Task, User
    from tests.conftest import TestingSessionLocal

    env: dict = {}
    with TestingSessionLocal() as s:
        def user(email: str, role: str, name: str) -> User:
            u = s.query(User).filter_by(email=email).one_or_none()
            if u is None:
                u = User(email=email, hashed_password=hash_password("pass"), name=name, role=role)
                s.add(u)
                s.flush()
            return u

        env["manager"] = user("p4mgr@test.jp", "PROJECT_MANAGER", "P4担当者A").id
        env["manager2"] = user("p4mgr2@test.jp", "PROJECT_MANAGER", "P4担当者B").id
        env["qm"] = user("p4qm@test.jp", "QUALITY_MANAGER", "P4品質").id
        env["viewer"] = user("p4viewer@test.jp", "VIEWER", "P4閲覧").id
        fw = user("p4fw@test.jp", "FIELD_WORKER", "P4現場")
        env["fw"] = fw.id

        dept = Department(name="P4部署")
        company_a = Company(name="P4会社A", is_partner=True)
        company_b = Company(name="P4会社B", is_partner=True)
        s.add_all([dept, company_a, company_b])
        s.flush()
        env["dept"] = dept.id
        env["company_a"] = company_a.id
        env["company_b"] = company_b.id

        def project(key: str, **over) -> None:
            base = dict(
                construction_number=f"{P4_TAG}-{key}",
                name=f"{P4_TAG} {key}",
                customer="P4顧客",
                status="遅延",
                area="P4エリア北",
                department_id=dept.id,
                manager_id=env["manager"],
                start_planned_at=date(2026, 3, 1),
                finish_planned_at=date(2026, 3, 31),
            )
            base.update(over)
            p = Project(**base)
            s.add(p)
            s.flush()
            env[key] = p.id

        # hit だけが全条件に一致する。他はどれか1条件だけ外れている。
        project("hit")
        project("status_ng", status="施工中")
        project("area_ng", area="P4エリア南")
        project("mgr_ng", manager_id=env["manager2"])
        project("dept_ng", department_id=None)
        project("period_ng", start_planned_at=date(2026, 8, 1),
                finish_planned_at=date(2026, 8, 31))
        project("company_ng")
        env["all"] = {env[k] for k in
                      ("hit", "status_ng", "area_ng", "mgr_ng", "dept_ng", "period_ng",
                       "company_ng")}

        # 会社は案件ではなく「工程の担当会社」。hit は同じ会社の工程を2件持つ。
        for key in ("hit", "status_ng", "area_ng", "mgr_ng", "dept_ng", "period_ng"):
            s.add_all([
                Task(project_id=env[key], name=f"{key} 工程1", company_id=company_a.id),
                Task(project_id=env[key], name=f"{key} 工程2", company_id=company_a.id),
            ])
        s.add(Task(project_id=env["hit"], name="hit 工程3", company_id=company_b.id))
        s.add(Task(project_id=env["company_ng"], name="company_ng 工程1",
                   company_id=company_b.id))

        # 期間検索の境界確認。片側だけ未設定の案件を必ず含める。
        periods = {
            "both": (date(2026, 5, 10), date(2026, 5, 20)),
            "no_start": (None, date(2026, 5, 20)),
            "no_finish": (date(2026, 5, 10), None),
            "no_period": (None, None),
            "before": (date(2026, 1, 1), date(2026, 1, 31)),
            "after": (date(2026, 9, 1), date(2026, 9, 30)),
        }
        for key, (start, finish) in periods.items():
            p = Project(construction_number=f"{P4_PERIOD_TAG}-{key}",
                        name=f"{P4_PERIOD_TAG} {key}", status="施工中",
                        start_planned_at=start, finish_planned_at=finish)
            s.add(p)
            s.flush()
            env[f"per_{key}"] = p.id
        env["periods"] = {env[f"per_{k}"] for k in periods}

        # 協力会社ロールは hit だけに割り当てる（0件と403を区別するため）
        if s.query(ProjectMember).filter_by(user_id=fw.id).one_or_none() is None:
            s.add(ProjectMember(project_id=env["hit"], user_id=fw.id, role="FIELD_WORKER"))
        s.commit()
    return env


def test_p4_project_lifecycle_create_list_detail_update(client, p4):
    """CREATE → 一覧 → 詳細 → UPDATE → 再取得 が一続きで通ること。"""
    head = _p4_head(client, "pm@test.jp")
    number = _p4_number()
    body = {
        "construction_number": number,
        "name": "ライフサイクル確認",
        "customer": "P4顧客",
        "status": "未着工",
        "area": "P4エリア北",
        "manager_id": p4["manager"],
        "start_planned_at": "2026-04-01",
        "finish_planned_at": "2026-04-30",
    }
    created = client.post("/api/projects", json=body, headers=head)
    assert created.status_code == 201, created.text
    new_id = created.json()["id"]

    # 一覧（検索API）。既定の並び順は登録が新しい順なので1ページ目の先頭に出る。
    listed = _p4_search(client, "pm@test.jp", "per_page=20&page=1")
    assert listed["sort"] == "recent"
    assert listed["items"][0]["id"] == new_id, "登録直後の案件が1ページ目の先頭に出ない"
    assert listed["page"] == 1 and listed["per_page"] == 20

    # 互換の一覧APIからも見える
    legacy = client.get("/api/projects", headers=head)
    assert new_id in [p["id"] for p in legacy.json()]

    # 詳細
    detail = client.get(f"/api/projects/{new_id}", headers=head)
    assert detail.status_code == 200, detail.text
    assert detail.json()["construction_number"] == number
    assert detail.json()["manager"] == "P4担当者A"

    # 更新
    updated = client.put(
        f"/api/projects/{new_id}",
        json={"name": "ライフサイクル確認（更新後）", "status": "施工中",
              "finish_planned_at": "2026-05-31", "actual_progress": 40},
        headers=head,
    )
    assert updated.status_code == 200, updated.text

    # 再取得で反映されている（画面の状態ではなく保存結果を確認する）
    again = client.get(f"/api/projects/{new_id}", headers=head).json()
    assert again["name"] == "ライフサイクル確認（更新後）"
    assert again["status"] == "施工中"
    assert again["finish_planned_at"] == "2026-05-31"
    assert again["actual_progress"] == 40
    assert again["start_planned_at"] == "2026-04-01", "更新していない項目が消えている"
    searched = _p4_search(client, "pm@test.jp", f"q={number}")
    assert [p["name"] for p in searched["items"]] == ["ライフサイクル確認（更新後）"]
    assert searched["total"] == 1


def test_p4_duplicate_construction_number_conflicts(client, p4):
    """工事番号の重複は、作成時も更新時も 409 で拒否されること。"""
    head = _p4_head(client, "pm@test.jp")
    first = _p4_number()
    r = client.post("/api/projects", json={"construction_number": first, "name": "重複元"},
                    headers=head)
    assert r.status_code == 201, r.text
    kept_id = r.json()["id"]

    dup = client.post("/api/projects", json={"construction_number": first, "name": "重複先"},
                      headers=head)
    assert dup.status_code == 409, dup.text
    assert _p4_search(client, "pm@test.jp", f"q={first}")["total"] == 1, "重複が保存されている"

    second = _p4_number()
    other = client.post("/api/projects", json={"construction_number": second, "name": "別案件"},
                        headers=head)
    assert other.status_code == 201, other.text
    other_id = other.json()["id"]

    conflict = client.put(f"/api/projects/{other_id}", json={"construction_number": first},
                          headers=head)
    assert conflict.status_code == 409, conflict.text
    assert client.get(f"/api/projects/{other_id}",
                      headers=head).json()["construction_number"] == second

    # 自分自身の工事番号での更新は重複扱いにしない
    same = client.put(f"/api/projects/{other_id}",
                      json={"construction_number": second, "name": "別案件（更新）"}, headers=head)
    assert same.status_code == 200, same.text
    assert client.get(f"/api/projects/{kept_id}", headers=head).status_code == 200


def test_p4_period_validation_rejects_reversed_dates(client, p4):
    """開始日が終了日より後なら 422。更新は保存後の値の組で判定すること。"""
    head = _p4_head(client, "pm@test.jp")
    bad = client.post(
        "/api/projects",
        json={"construction_number": _p4_number(), "name": "逆転期間",
              "start_planned_at": "2026-06-30", "finish_planned_at": "2026-06-01"},
        headers=head,
    )
    assert bad.status_code == 422, bad.text

    number = _p4_number()
    ok = client.post(
        "/api/projects",
        json={"construction_number": number, "name": "期間更新",
              "start_planned_at": "2026-06-01", "finish_planned_at": "2026-06-30"},
        headers=head,
    )
    assert ok.status_code == 201, ok.text
    pid = ok.json()["id"]

    # 片側だけ更新しても、保存済みのもう片方と突き合わせて判定される
    late_start = client.put(f"/api/projects/{pid}", json={"start_planned_at": "2026-07-15"},
                            headers=head)
    assert late_start.status_code == 422, late_start.text
    early_finish = client.put(f"/api/projects/{pid}", json={"finish_planned_at": "2026-05-01"},
                              headers=head)
    assert early_finish.status_code == 422, early_finish.text
    kept = client.get(f"/api/projects/{pid}", headers=head).json()
    assert kept["start_planned_at"] == "2026-06-01" and kept["finish_planned_at"] == "2026-06-30"

    # 同日は許可（1日で終わる工事がある）
    same_day = client.put(f"/api/projects/{pid}",
                          json={"start_planned_at": "2026-06-30"}, headers=head)
    assert same_day.status_code == 200, same_day.text


def test_p4_unique_violation_maps_to_conflict(client, p4, monkeypatch):
    """事前チェックをすり抜けても、DBの一意制約違反は 500 ではなく 409 になること。"""
    head = _p4_head(client, "pm@test.jp")
    number = _p4_number()
    seed = client.post("/api/projects", json={"construction_number": number, "name": "一意制約元"},
                       headers=head)
    assert seed.status_code == 201, seed.text

    other_number = _p4_number()
    other = client.post("/api/projects",
                        json={"construction_number": other_number, "name": "一意制約先"},
                        headers=head)
    assert other.status_code == 201, other.text
    other_id = other.json()["id"]

    # 事前チェックを無効化して IntegrityError の経路だけを通す
    monkeypatch.setattr("app.api.projects._check_number_unique", lambda *a, **k: None)

    dup = client.post("/api/projects", json={"construction_number": number, "name": "衝突"},
                      headers=head)
    assert dup.status_code == 409, dup.text

    conflict = client.put(f"/api/projects/{other_id}", json={"construction_number": number},
                          headers=head)
    assert conflict.status_code == 409, conflict.text

    monkeypatch.undo()
    # ロールバックされ、既存データは壊れていない
    assert client.get(f"/api/projects/{other_id}",
                      headers=head).json()["construction_number"] == other_number
    assert _p4_search(client, "pm@test.jp", f"q={number}")["total"] == 1


def test_p4_search_combines_all_filters_with_and(client, p4):
    """全条件を同時指定するとANDで効き、条件を1つ外すとその分だけ増えること。"""
    base = {
        "q": P4_TAG,
        "statuses": "遅延",
        "delayed_only": "true",
        "manager_ids": str(p4["manager"]),
        "department_ids": str(p4["dept"]),
        "company_ids": str(p4["company_a"]),
        "areas": "P4エリア北",
        "date_from": "2026-03-01",
        "date_to": "2026-03-31",
        "per_page": "50",
    }

    def qs(*drop: str) -> str:
        return "&".join(f"{k}={v}" for k, v in base.items() if k not in drop)

    assert set(_p4_ids(client, "admin@test.jp", qs())) == {p4["hit"]}

    # 条件を1つ外すと、その条件で落ちていた案件だけが戻る
    assert set(_p4_ids(client, "admin@test.jp", qs("areas"))) == {p4["hit"], p4["area_ng"]}
    assert set(_p4_ids(client, "admin@test.jp", qs("manager_ids"))) == {p4["hit"], p4["mgr_ng"]}
    assert set(_p4_ids(client, "admin@test.jp", qs("department_ids"))) == {p4["hit"], p4["dept_ng"]}
    assert set(_p4_ids(client, "admin@test.jp", qs("company_ids"))) == {p4["hit"], p4["company_ng"]}
    assert set(_p4_ids(client, "admin@test.jp", qs("date_to"))) == {p4["hit"], p4["period_ng"]}
    assert set(_p4_ids(client, "admin@test.jp", qs("statuses", "delayed_only"))) == {
        p4["hit"], p4["status_ng"]}

    # キーワードは工事名・工事番号・顧客・責任者名のいずれかに当たる
    assert set(_p4_ids(client, "admin@test.jp", f"q={P4_TAG}&per_page=50")) == p4["all"]
    assert set(_p4_ids(client, "admin@test.jp", "q=P4担当者B&per_page=50")) == {p4["mgr_ng"]}
    assert _p4_search(client, "admin@test.jp", "q=該当しないキーワード")["total"] == 0


def test_p4_search_by_task_company_without_duplicates(client, p4):
    """会社は工程の担当会社で判定し、同じ会社の工程が複数でも案件が重複しないこと。"""
    both = f"{p4['company_a']},{p4['company_b']}"
    ids = _p4_ids(client, "admin@test.jp", f"q={P4_TAG}&company_ids={both}&per_page=50")
    assert len(ids) == len(set(ids)), f"同じ案件が複数行で返っている: {ids}"
    # hit は会社Aの工程を2件、会社Bの工程を1件持つが1件として数える
    assert ids.count(p4["hit"]) == 1
    assert set(ids) == p4["all"]
    assert _p4_search(client, "admin@test.jp",
                      f"q={P4_TAG}&company_ids={both}&per_page=50")["total"] == len(p4["all"])

    only_b = _p4_ids(client, "admin@test.jp",
                     f"q={P4_TAG}&company_ids={p4['company_b']}&per_page=50")
    assert set(only_b) == {p4["hit"], p4["company_ng"]}
    assert len(only_b) == 2


def test_p4_search_period_treats_open_ended_as_overlapping(client, p4):
    """期間は「予定期間が範囲と重なる案件」。片側未設定はその向きに制限しない。"""
    tag = f"q={P4_PERIOD_TAG}&per_page=50"

    def ids(extra: str) -> set:
        return set(_p4_ids(client, "admin@test.jp", f"{tag}&{extra}"))

    # 境界日ちょうどを含む（5/20 は both の完了予定日、no_start の完了予定日）
    assert ids("date_from=2026-05-20&date_to=2026-05-20") == {
        p4["per_both"], p4["per_no_start"], p4["per_no_finish"], p4["per_no_period"]}

    # 1日ずらすと、終了予定日で外れる案件が落ちる
    assert ids("date_from=2026-05-21&date_to=2026-05-21") == {
        p4["per_no_finish"], p4["per_no_period"]}

    # 開始予定日の境界。5/10 は both / no_finish の着工予定日
    assert ids("date_to=2026-05-10&date_from=2026-05-10") == {
        p4["per_both"], p4["per_no_start"], p4["per_no_finish"], p4["per_no_period"]}
    assert ids("date_to=2026-05-09&date_from=2026-05-09") == {
        p4["per_no_start"], p4["per_no_period"]}

    # 片側だけの指定は、その向きだけ制限する
    assert ids("date_from=2026-06-01") == {
        p4["per_no_finish"], p4["per_no_period"], p4["per_after"]}
    assert ids("date_to=2026-02-01") == {
        p4["per_no_start"], p4["per_no_period"], p4["per_before"]}

    # 期間未指定なら全件
    assert ids("") == p4["periods"]


def test_p4_total_is_counted_after_scope_and_pages_sum_to_total(client, p4):
    """total は権限・条件を適用したあとの件数で、全ページの合計と一致すること。"""
    qs = f"q={P4_TAG}"
    admin_ids, admin_total, admin_pages = _p4_walk_pages(client, "admin@test.jp", qs, per_page=2)
    assert admin_total == len(p4["all"])
    assert len(admin_ids) == admin_total, "全ページの items 合計が total と一致しない"
    assert len(set(admin_ids)) == admin_total, "ページ間で同じ案件が重複している"
    assert set(admin_ids) == p4["all"]
    assert admin_pages == (admin_total + 1) // 2

    # 協力会社ロールは割当案件だけ。total もスコープ適用後の値になる。
    fw_ids, fw_total, _ = _p4_walk_pages(client, "p4fw@test.jp", qs, per_page=2)
    assert fw_ids == [p4["hit"]]
    assert fw_total == 1, "権限適用前の件数を total に返している"
    assert fw_total < admin_total

    # 1ページに収めても total は変わらない
    single = _p4_search(client, "admin@test.jp", f"{qs}&per_page=50")
    assert single["total"] == admin_total and single["pages"] == 1
    assert len(single["items"]) == admin_total


def test_p4_filter_options_cover_whole_scope_not_current_page(client, p4):
    """絞り込みの選択肢は表示中のページではなく、権限範囲の全案件から作ること。"""
    head = _p4_head(client, "admin@test.jp")
    page = _p4_search(client, "admin@test.jp", f"q={P4_TAG}&per_page=1&page=1")
    assert len(page["items"]) == 1 and page["total"] > 1

    r = client.get("/api/projects/filter-options", headers=head)
    assert r.status_code == 200, r.text
    opts = r.json()
    areas = set(opts["areas"])
    assert {"P4エリア北", "P4エリア南"} <= areas, "1ページ目に無いエリアが選択肢から落ちている"
    assert {"P4担当者A", "P4担当者B"} <= {m["name"] for m in opts["managers"]}
    assert {"P4会社A", "P4会社B"} <= {c["name"] for c in opts["companies"]}
    assert "P4部署" in {d["name"] for d in opts["departments"]}
    assert {"遅延", "施工中"} <= set(opts["statuses"])
    assert len({d["id"] for d in opts["departments"]}) == len(opts["departments"])

    # 権限範囲が狭ければ選択肢も狭まる（割当案件は hit のみ）
    fw = client.get("/api/projects/filter-options",
                    headers=_p4_head(client, "p4fw@test.jp")).json()
    assert fw["areas"] == ["P4エリア北"]
    assert [m["name"] for m in fw["managers"]] == ["P4担当者A"]
    assert {c["name"] for c in fw["companies"]} == {"P4会社A", "P4会社B"}
    assert fw["statuses"] == ["遅延"]


def test_p4_static_routes_do_not_collide_with_project_id(client, p4):
    """/search と /filter-options が /{project_id} に吸われていないこと。"""
    head = _p4_head(client, "admin@test.jp")

    search = client.get("/api/projects/search", headers=head)
    assert search.status_code == 200, search.text
    assert set(search.json()) >= {"items", "total", "page", "per_page", "pages", "sort"}
    assert "construction_number" not in search.json(), "検索が案件詳細として解釈されている"

    options = client.get("/api/projects/filter-options", headers=head)
    assert options.status_code == 200, options.text
    assert set(options.json()) >= {"statuses", "areas", "departments", "managers", "companies"}

    detail = client.get(f"/api/projects/{p4['hit']}", headers=head)
    assert detail.status_code == 200, detail.text
    assert detail.json()["id"] == p4["hit"]

    missing = client.get("/api/projects/99999999", headers=head)
    assert missing.status_code == 404, missing.text

    # 数値に変換できないパスは 422（= 動的ルートの側で弾かれる）
    assert client.get("/api/projects/not-a-number", headers=head).status_code == 422


def test_p4_update_allowed_for_admin_and_project_manager(client, p4):
    """ADMIN と PROJECT_MANAGER は案件を更新できること。"""
    for email, label in (("admin@test.jp", "管理者更新"), ("pm@test.jp", "PM更新")):
        head = _p4_head(client, email)
        created = client.post("/api/projects",
                              json={"construction_number": _p4_number(), "name": "権限確認"},
                              headers=head)
        assert created.status_code == 201, created.text
        pid = created.json()["id"]
        r = client.put(f"/api/projects/{pid}", json={"name": label}, headers=head)
        assert r.status_code == 200, r.text
        assert client.get(f"/api/projects/{pid}", headers=head).json()["name"] == label


def test_p4_update_forbidden_for_other_roles(client, p4):
    """QUALITY_MANAGER・FIELD_WORKER・VIEWER は API 側で 403 になること。"""
    admin = _p4_head(client, "admin@test.jp")
    before = client.get(f"/api/projects/{p4['hit']}", headers=admin).json()["name"]

    for email in ("p4qm@test.jp", "p4fw@test.jp", "p4viewer@test.jp"):
        head = _p4_head(client, email)
        upd = client.put(f"/api/projects/{p4['hit']}", json={"name": "権限外の更新"}, headers=head)
        assert upd.status_code == 403, f"{email} の更新が通ってしまった: {upd.text}"
        create = client.post("/api/projects",
                             json={"construction_number": _p4_number(), "name": "権限外の登録"},
                             headers=head)
        assert create.status_code == 403, f"{email} の登録が通ってしまった: {create.text}"

    assert client.get(f"/api/projects/{p4['hit']}", headers=admin).json()["name"] == before


def test_p4_field_worker_cannot_reach_unassigned_project(client, p4):
    """割当外の案件は、取得も更新もできないこと。"""
    head = _p4_head(client, "p4fw@test.jp")
    assert client.get(f"/api/projects/{p4['hit']}", headers=head).status_code == 200

    unassigned = p4["area_ng"]
    assert client.get(f"/api/projects/{unassigned}", headers=head).status_code == 403
    assert client.put(f"/api/projects/{unassigned}", json={"name": "割当外の更新"},
                      headers=head).status_code == 403
    # 検索結果にも出ない
    assert unassigned not in _p4_ids(client, "p4fw@test.jp", f"q={P4_TAG}&per_page=50")


def test_p4_empty_result_is_distinguishable_from_forbidden(client, p4):
    """0件（200 + total 0）と権限なし（403）を取り違えないこと。"""
    empty = _p4_search(client, "p4fw@test.jp", "q=該当しないキーワード")
    assert empty["total"] == 0 and empty["items"] == [] and empty["pages"] == 1

    # 権限範囲は空ではない
    assert _p4_search(client, "p4fw@test.jp", f"q={P4_TAG}")["total"] == 1

    forbidden = client.get(f"/api/projects/{p4['area_ng']}",
                           headers=_p4_head(client, "p4fw@test.jp"))
    assert forbidden.status_code == 403

    # 割当が1件も無い協力会社は、403 ではなく 0件の 200 になる
    from app.core.security import hash_password
    from app.models import User
    from tests.conftest import TestingSessionLocal

    with TestingSessionLocal() as s:
        if s.query(User).filter_by(email="p4fw0@test.jp").one_or_none() is None:
            s.add(User(email="p4fw0@test.jp", hashed_password=hash_password("pass"),
                       name="P4未割当", role="FIELD_WORKER"))
            s.commit()
    none_assigned = _p4_search(client, "p4fw0@test.jp", "per_page=50")
    assert none_assigned["total"] == 0 and none_assigned["items"] == []
    opts = client.get("/api/projects/filter-options",
                      headers=_p4_head(client, "p4fw0@test.jp"))
    assert opts.status_code == 200
    assert opts.json() == {"statuses": [], "areas": [], "departments": [], "managers": [],
                           "companies": []}


def test_p4_upload_rejects_missing_project(client, p4):
    """案件が決まっていない写真アップロードは、既定値で登録されず拒否されること。

    画面側は案件未選択ならアップロード自体をさせないが、API単体でも
    存在しない案件ID（0 など）を受け付けないことを固定する。
    """
    from app.models import Photo
    from tests.conftest import TestingSessionLocal

    head = _p4_head(client, "pm@test.jp")
    with TestingSessionLocal() as s:
        before = s.query(Photo).count()

    files = {"file": ("t.jpg", b"not-an-image", "image/jpeg")}
    zero = client.post("/api/photos", data={"project_id": "0"}, files=files, headers=head)
    assert zero.status_code == 404, zero.text
    missing = client.post("/api/photos", data={"project_id": "99999999"}, files=files, headers=head)
    assert missing.status_code == 404, missing.text
    # 案件IDそのものが無いリクエストは通さない
    no_project = client.post("/api/photos", data={}, files=files, headers=head)
    assert no_project.status_code in (400, 422), no_project.text

    with TestingSessionLocal() as s:
        assert s.query(Photo).count() == before, "拒否したのに写真が作られている"


def test_p4_test_record_requires_accessible_project(client, p4):
    """試験記録も、案件スコープ外・存在しない案件には登録できないこと。"""
    from app.models import TestRecord
    from tests.conftest import TestingSessionLocal

    with TestingSessionLocal() as s:
        before = s.query(TestRecord).count()

    body = {"project_id": p4["hit"], "test_type": "光損失測定", "judge": "合格"}
    ok_res = client.post("/api/test-records", json=body, headers=_p4_head(client, "pm@test.jp"))
    assert ok_res.status_code in (200, 201), ok_res.text

    # 割当外の案件（協力会社ロール）
    ng = client.post(
        "/api/test-records",
        json={**body, "project_id": p4["area_ng"]},
        headers=_p4_head(client, "p4fw@test.jp"),
    )
    assert ng.status_code == 403, ng.text

    missing = client.post(
        "/api/test-records", json={**body, "project_id": 99999999},
        headers=_p4_head(client, "pm@test.jp"),
    )
    assert missing.status_code == 404, missing.text

    with TestingSessionLocal() as s:
        assert s.query(TestRecord).count() == before + 1, "拒否した分まで登録されている"


# ---------------------------------------------------------------------------
# Phase 5: 帳票（画面・プレビュー・PDF・Excel・CSV が同じ内容であること）
# ---------------------------------------------------------------------------

def _p5_project_with_tasks(client):
    """工程を持つ検証用の案件を作る。Seed の固定IDには依存しない。"""
    from datetime import datetime, timedelta, timezone

    from app.models import Task
    from tests.conftest import TestingSessionLocal

    head = _p4_head(client, "pm@test.jp")
    number = _p4_number("P5RPT")
    created = client.post("/api/projects", json={"construction_number": number, "name": "帳票の一致確認"},
                          headers=head)
    assert created.status_code == 201, created.text
    pid = created.json()["id"]

    jst = timezone(timedelta(hours=9))
    base = datetime(2026, 4, 1, 0, 0, tzinfo=jst)
    with TestingSessionLocal() as s:
        # 親を持たない工程（WBS にドットが無い）だけの案件でも帳票へ載ること
        s.add(Task(project_id=pid, wbs_code="1", name="現場確認",
                   planned_start_at=base, planned_finish_at=base + timedelta(days=1),
                   actual_start_at=base, actual_finish_at=base + timedelta(days=1),
                   planned_workers=3, actual_workers=4, actual_progress=100, status="完了",
                   notes="立会あり"))
        s.add(Task(project_id=pid, wbs_code="2", name="現地確認",
                   planned_start_at=base + timedelta(days=1), planned_finish_at=base + timedelta(days=2),
                   planned_workers=2, actual_workers=0, actual_progress=30, status="施工中"))
        s.commit()
    return pid, number, head


def _p5_preview(client, pid, head):
    r = client.get(f"/api/reports/construction-management/preview?project_id={pid}", headers=head)
    assert r.status_code == 200, r.text
    return r.json()


def test_p5_report_includes_tasks_without_parent(client, p4):
    """親子を作っていない案件でも、工程が帳票に載ること（0件にならない）。"""
    pid, number, head = _p5_project_with_tasks(client)
    body = _p5_preview(client, pid, head)
    assert body["row_count"] == 2, body["rows"]
    assert [r[1] for r in body["rows"]] == ["現場確認", "現地確認"], "WBS順で工程が並ばない"
    assert body["construction_number"] == number
    assert body["source"] == "tasks"
    # 画面が固定行を作らないよう、行はすべて工程由来であること
    names = {r[1] for r in body["rows"]}
    assert not names & {"光ケーブル敷設", "クロージャ設置", "光ファイバ融着", "接続損失測定", "ONU設置"}


def test_p5_report_screen_pdf_xlsx_csv_match(client, p4):
    """プレビュー（画面）・PDF・Excel・CSV の件数と内容が一致すること。"""
    import csv as _csv
    import io
    import re

    from openpyxl import load_workbook
    from pypdf import PdfReader

    pid, _number, head = _p5_project_with_tasks(client)
    body = _p5_preview(client, pid, head)
    rows = body["rows"]
    assert len(rows) == 2

    # --- Excel: セルを読んで突き合わせる ---
    xlsx = client.get(f"/api/reports/construction-management?project_id={pid}&format=xlsx", headers=head)
    assert xlsx.status_code == 200, xlsx.text
    wb = load_workbook(io.BytesIO(xlsx.content))
    ws = wb.active
    cells = [[("" if c is None else str(c)) for c in row] for row in ws.iter_rows(values_only=True)]
    header_at = next(i for i, r in enumerate(cells) if r[:len(body["columns"])] == body["columns"])
    xlsx_rows = [r[:len(body["columns"])] for r in cells[header_at + 1:] if any(v for v in r)]
    assert xlsx_rows == rows, f"Excel が画面と一致しない: {xlsx_rows}"

    # --- CSV: パースして突き合わせる ---
    csv_res = client.get(f"/api/reports/construction-management?project_id={pid}&format=csv", headers=head)
    assert csv_res.status_code == 200, csv_res.text
    text = csv_res.content.decode("utf-8-sig")
    parsed = list(_csv.reader(io.StringIO(text)))
    c_header = next(i for i, r in enumerate(parsed) if r == body["columns"])
    csv_rows = [r for r in parsed[c_header + 1:] if any(v for v in r)]
    assert csv_rows == rows, f"CSV が画面と一致しない: {csv_rows}"

    # --- PDF: テキストを抽出して突き合わせる ---
    pdf = client.get(f"/api/reports/construction-management?project_id={pid}&format=pdf", headers=head)
    assert pdf.status_code == 200, pdf.text
    text = "".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(pdf.content)).pages)
    flat = re.sub(r"\s+", "", text)
    for row in rows:
        for value in row:
            if value in ("—", ""):
                continue
            assert re.sub(r"\s+", "", value) in flat, f"PDFに {value} が無い"
    assert "施工管理表" in flat


def test_p5_report_empty_project_has_no_rows(client, p4):
    """工程0件の案件では、固定行を出さずどの形式も0件になること。"""
    import csv as _csv
    import io

    from openpyxl import load_workbook

    head = _p4_head(client, "pm@test.jp")
    created = client.post("/api/projects",
                          json={"construction_number": _p4_number("P5EMP"), "name": "工程なしの帳票"},
                          headers=head)
    assert created.status_code == 201, created.text
    pid = created.json()["id"]

    body = _p5_preview(client, pid, head)
    assert body["row_count"] == 0 and body["rows"] == []

    xlsx = client.get(f"/api/reports/construction-management?project_id={pid}&format=xlsx", headers=head)
    wb = load_workbook(io.BytesIO(xlsx.content))
    cells = [[("" if c is None else str(c)) for c in r] for r in wb.active.iter_rows(values_only=True)]
    header_at = next(i for i, r in enumerate(cells) if r[:len(body["columns"])] == body["columns"])
    assert [r for r in cells[header_at + 1:] if any(v for v in r)] == []

    csv_res = client.get(f"/api/reports/construction-management?project_id={pid}&format=csv", headers=head)
    parsed = list(_csv.reader(io.StringIO(csv_res.content.decode("utf-8-sig"))))
    c_header = next(i for i, r in enumerate(parsed) if r == body["columns"])
    assert [r for r in parsed[c_header + 1:] if any(v for v in r)] == []


def test_p5_report_meta_comes_from_real_data(client, p4):
    """工事名・工事番号・出力日時・出力者・件数が実データから作られること。"""
    from datetime import datetime, timezone

    pid, number, head = _p5_project_with_tasks(client)
    body = _p5_preview(client, pid, head)
    meta = {m["label"]: m["value"] for m in body["meta"]}
    assert meta["工事番号"] == number
    assert meta["工事名"] == "帳票の一致確認"
    assert meta["工程件数"] == "2 件"
    assert meta["出力者"] == "PM"  # ログイン中の利用者
    assert meta["出力日時"].startswith(str(datetime.now(timezone.utc).year))
    assert "熊本中央局" not in str(meta)


def test_p5_report_respects_permissions(client, p4):
    """帳票も案件スコープと権限に従うこと。"""
    pid, _n, _h = _p5_project_with_tasks(client)
    for path in (f"/api/reports/construction-management/preview?project_id={pid}",
                 f"/api/reports/construction-management?project_id={pid}&format=csv"):
        r = client.get(path, headers=_p4_head(client, "p4fw@test.jp"))
        assert r.status_code == 403, f"{path} が割当外へ漏れている: {r.status_code}"
    bad = client.get(f"/api/reports/construction-management?project_id={pid}&format=docx",
                     headers=_p4_head(client, "pm@test.jp"))
    assert bad.status_code == 422
    missing = client.get(f"/api/reports/unknown-report/preview?project_id={pid}",
                         headers=_p4_head(client, "pm@test.jp"))
    assert missing.status_code == 404


def test_p5_milestone_options_allow_projects_without_milestones(client, p4):
    """マイルストーンが1件も無い案件でも、登録先として選べること。

    録画で「案件欄が『選択してください』のまま候補も出ない」状態になっていた原因は、
    選択肢を既存のマイルストーンから作っていたこと。登録できる案件から作る。
    """
    head = _p4_head(client, "pm@test.jp")
    created = client.post("/api/projects",
                          json={"construction_number": _p4_number("P5MS"), "name": "マイルストーン未登録の案件"},
                          headers=head)
    assert created.status_code == 201, created.text
    pid = created.json()["id"]

    # 案件配下ルート相当（project_id 指定）
    scoped = client.get(f"/api/schedule/milestones/options?project_id={pid}", headers=head).json()
    assert [p["id"] for p in scoped["projects"]] == [pid], "自分の案件が候補に出ない"
    assert scoped["milestone_types"], "マイルストーン種別マスタが取得できない"

    # 横断画面（project_id なし）でも権限範囲の案件がすべて選べる
    across = client.get("/api/schedule/milestones/options", headers=head).json()
    assert pid in [p["id"] for p in across["projects"]]
    assert len(across["projects"]) >= len(scoped["projects"])

    # 権限が無い案件は候補に出ない
    fw = client.get("/api/schedule/milestones/options", headers=_p4_head(client, "p4fw@test.jp")).json()
    assert pid not in [p["id"] for p in fw["projects"]]
    assert [p["id"] for p in fw["projects"]] == [p4["hit"]]


def test_p5_milestone_related_tasks_come_from_project(client, p4):
    """関連工程は、その案件の工程から選べること（既存の紐付けに限定しない）。"""
    from app.models import Task
    from tests.conftest import TestingSessionLocal

    head = _p4_head(client, "pm@test.jp")
    created = client.post("/api/projects",
                          json={"construction_number": _p4_number("P5MST"), "name": "関連工程の選択肢"},
                          headers=head)
    pid = created.json()["id"]
    with TestingSessionLocal() as s:
        s.add(Task(project_id=pid, wbs_code="1", name="関連工程の候補"))
        s.commit()

    opts = client.get(f"/api/schedule/milestones/options?project_id={pid}", headers=head).json()
    names = [t["name"] for t in opts["related_tasks"]]
    assert "関連工程の候補" in names, names
    assert all(t["project_id"] == pid for t in opts["related_tasks"]), "他案件の工程が混ざっている"


def test_p5_milestone_create_from_project_route(client, p4):
    """案件配下から登録したマイルストーンが、その案件で取得できること。"""
    head = _p4_head(client, "pm@test.jp")
    created = client.post("/api/projects",
                          json={"construction_number": _p4_number("P5MSC"), "name": "マイルストーン登録"},
                          headers=head)
    pid = created.json()["id"]
    opts = client.get(f"/api/schedule/milestones/options?project_id={pid}", headers=head).json()
    type_id = opts["milestone_types"][0]["id"]

    r = client.post("/api/schedule/milestones", json={
        "project_id": pid, "milestone_type_id": type_id, "name": "着工",
        "planned_at": "2026-05-01T00:00:00+09:00", "status": "予定", "schedule_precision": "day",
    }, headers=head)
    assert r.status_code in (200, 201), r.text

    listed = client.get(f"/api/schedule/milestones?project_id={pid}&limit=100", headers=head).json()
    assert [m["name"] for m in listed["milestones"]] == ["着工"]
    assert listed["milestones"][0]["project_id"] == pid


def test_p5_system_info_shows_revision_without_secrets(client, p4):
    """稼働中のコード世代と接続状態を返し、秘密情報は返さないこと。"""
    admin = _p4_head(client, "admin@test.jp")
    r = client.get("/api/system/info", headers=admin)
    assert r.status_code == 200, r.text
    body = r.json()
    for key in ("environment", "api_version", "backend_commit", "backend_built_at",
                "alembic_revision", "database", "storage", "ai_service", "server_time"):
        assert key in body, f"{key} が返っていない"
    assert body["database"] == "ok"
    # 秘密情報・接続先を返さない
    dumped = str(body).lower()
    for leaked in ("password", "secret", "token", "postgresql://", "sqlite://", "jwt", "key="):
        assert leaked not in dumped, f"{leaked} が含まれている"

    # 管理者以外は参照できない
    for email in ("pm@test.jp", "p4qm@test.jp", "p4viewer@test.jp", "p4fw@test.jp"):
        assert client.get("/api/system/info", headers=_p4_head(client, email)).status_code == 403

    # 稼働確認だけは全ロールで取れる（秘密情報なし）
    ping = client.get("/api/system/ping", headers=_p4_head(client, "p4viewer@test.jp"))
    assert ping.status_code == 200
    assert set(ping.json()) == {"environment", "api_version", "backend_commit"}


def test_p5_health_exposes_code_generation(client):
    """/health からも稼働中のコード世代が分かること（認証不要・秘密情報なし）。"""
    r = client.get("/health")
    assert r.status_code == 200
    body = r.json()
    assert body["status"] == "ok"
    for key in ("version", "environment", "backend_commit", "backend_built_at"):
        assert key in body
    assert "password" not in str(body).lower()
